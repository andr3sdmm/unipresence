"""
Tests for UniPresence v2.

They run against a real Postgres database, because the behaviour that matters
most (two people racing for one carne, one transaction per request) is the
database's behaviour. Point TEST_DATABASE_URL at an empty database you can wipe:

    TEST_DATABASE_URL="postgresql://..." pytest -q

Every test run drops and recreates the unipresence_v2 schema in that database.
Never point it at the production database.
"""

import csv
import io
import os
import re

import psycopg
import pytest

TEST_DB = os.environ.get("TEST_DATABASE_URL")
if not TEST_DB:
    pytest.skip("TEST_DATABASE_URL is not set", allow_module_level=True)
os.environ["DATABASE_URL"] = TEST_DB

from fastapi.testclient import TestClient  # noqa: E402

import main  # noqa: E402

IPHONE = ("Mozilla/5.0 (iPhone; CPU iPhone OS 17_5 like Mac OS X) AppleWebKit/605.1.15 "
          "(KHTML, like Gecko) Version/17.5 Mobile/15E148 Safari/604.1")
ANDROID_CHROME = ("Mozilla/5.0 (Linux; Android 14; Pixel 8) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/126.0.0.0 Mobile Safari/537.36")
INSTAGRAM = IPHONE + " Instagram 330.0.0"


@pytest.fixture(autouse=True)
def fresh_schema():
    with psycopg.connect(TEST_DB, autocommit=True) as conn:
        conn.execute(f"DROP SCHEMA IF EXISTS {main.SCHEMA} CASCADE")
    main.init_db()
    yield


def sql(query, params=()):
    with psycopg.connect(TEST_DB, autocommit=True) as conn:
        conn.execute(f"SET search_path TO {main.SCHEMA}")
        cur = conn.execute(query, params)
        return cur.fetchall() if cur.description else None


def teacher(email="prof@uni.edu", name="Prof. Rojas"):
    c = TestClient(main.app)
    r = c.post("/signup", data={"name": name, "email": email, "password": "clave-segura-1",
                                "institution": "Universidad"}, follow_redirects=False)
    assert r.status_code == 303
    return c


def student(ua=IPHONE):
    return TestClient(main.app, headers={"user-agent": ua})


def create_course(t, name="Genética G1", roster=""):
    r = t.post("/course/new", data={"name": name, "roster_text": roster}, follow_redirects=False)
    return int(re.search(r"/course/(\d+)", r.headers["location"]).group(1))


def start(t, course_id):
    r = t.post(f"/course/{course_id}/start", follow_redirects=False)
    return int(re.search(r"/session/(\d+)/live", r.headers["location"]).group(1))


def fresh_nonce(t, sid):
    assert t.get(f"/session/{sid}/qr").status_code == 200
    return sql("SELECT nonce FROM nonces WHERE session_id = %s ORDER BY expires_at DESC LIMIT 1",
               (sid,))[0][0]


def ticket_of(page):
    return re.search(r'name="ticket" value="([^"]+)"', page).group(1)


def scan(s, t, sid):
    return s.get(f"/checkin?s={sid}&n={fresh_nonce(t, sid)}")


ROSTER = "María Fernanda Rojas\t2021104567\nPérez, Juan\t2021104568\nValentina Castro\t2021104569"


def test_roster_paste_keeps_names_with_commas():
    t = teacher()
    cid = create_course(t, roster=ROSTER)
    names = {r[0] for r in sql("SELECT name FROM students")}
    assert names == {"María Fernanda Rojas", "Pérez, Juan", "Valentina Castro"}
    assert sql("SELECT has_roster FROM courses WHERE id = %s", (cid,))[0][0] is True


def test_first_checkin_then_one_tap():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    page = scan(s, t, sid).text
    assert "Primera vez en este celular" in page
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(page), "carne": "2021104567"})
    assert "Asistencia recibida" in r.text and "María Fernanda Rojas" in r.text
    assert re.search(r"C\d+-\d{4}-[A-Z2-9]{4}", r.text)

    # Next class: one tap, no typing.
    sid2 = start(t, sql("SELECT course_id FROM sessions WHERE id = %s", (sid,))[0][0])
    page2 = scan(s, t, sid2).text
    assert "Aceptar" in page2 and "Primera vez" not in page2 and "María Fernanda Rojas" in page2
    assert "Asistencia recibida" in s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(page2)}).text


def test_second_scan_same_class_shows_existing_receipt():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "2021104567"})
    assert "Ya tenías asistencia" in scan(s, t, sid).text
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 1


def test_expired_qr_is_rejected_and_counted():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    n = fresh_nonce(t, sid)
    sql("UPDATE nonces SET expires_at = now() - interval '1 second' WHERE nonce = %s", (n,))
    assert "Vuelve a escanear" in student().get(f"/checkin?s={sid}&n={n}").text
    assert sql("SELECT COUNT(*) FROM events WHERE kind = 'qr_vencido'")[0][0] == 1


def test_ticket_cannot_be_used_twice():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    tk = ticket_of(scan(s, t, sid).text)
    s.post("/checkin/accept", data={"acepto": "1", "ticket": tk, "carne": "2021104567"})
    assert "ya se usó" in s.post("/checkin/accept", data={"acepto": "1", "ticket": tk, "carne": "2021104567"}).text


def test_carne_not_on_list_asks_for_name_then_goes_to_review():
    """Never lock a present student out: not refused, marked as pending and flagged."""
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    tk = ticket_of(scan(s, t, sid).text)
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": tk, "carne": "9999999"})
    assert "No encontramos" in r.text and 'name="nombre"' in r.text
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 0
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": tk, "carne": "9999999", "nombre": "Eva Marin"})
    assert "Asistencia recibida" in r.text
    assert sql("SELECT en.status FROM enrollments en JOIN students st ON st.id = en.student_id "
               "WHERE st.carne = '9999999'")[0][0] == "pendiente"
    assert sql("SELECT review FROM events WHERE kind = 'carne_no_en_lista'")[0][0] is True
    labels = [a["label"] for a in t.get(f"/session/{sid}/state").json()["alerts"]]
    assert main.EVENT_LABELS["carne_no_en_lista"] in labels


def test_brightspace_list_matches_the_username_students_know():
    """Real case, 19 Oct 2026: OrgDefinedId is a number students do not know."""
    roster = ("OrgDefinedId,Last Name,First Name,Email\n"
              "170001,GOMEZ DIAZ,ANA,agomez@universidad.edu.co\n"
              "170002,PEREZ ROA,LUIS,lperez@universidad.edu.co\n")
    t = teacher()
    sid = start(t, create_course(t, roster=roster))
    assert sql("SELECT username FROM students ORDER BY carne") == [("AGOMEZ",), ("LPEREZ",)]
    assert not sql("SELECT 1 FROM students WHERE username LIKE '%%@%%'")  # the address is never kept
    ana = student()
    r = ana.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(ana, t, sid).text),
                                          "carne": "agomez"})
    assert "Asistencia recibida" in r.text and "Ana Gomez Diaz" in r.text
    luis = student(ANDROID_CHROME)
    r = luis.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(luis, t, sid).text),
                                           "carne": "LPerez@universidad.edu.co"})
    assert "Luis Perez Roa" in r.text
    assert sql("SELECT COUNT(*) FROM events WHERE kind = 'carne_no_en_lista'")[0][0] == 0


def test_national_id_column_is_never_used_as_key():
    t = teacher()
    create_course(t, roster="Nombre,Cedula\nAna Gomez,1045678901\n")
    assert sql("SELECT COUNT(*) FROM students")[0][0] == 0


def test_known_browser_cannot_mark_for_a_friend():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "2021104567"})
    sid2 = start(t, sql("SELECT course_id FROM sessions WHERE id = %s", (sid,))[0][0])
    # The form for a known browser has no carne field; a crafted one is ignored.
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid2).text),
                                        "carne": "2021104568"})
    assert "María Fernanda Rojas" in r.text
    marked = {row[0] for row in sql("SELECT st.carne FROM attendance a JOIN students st "
                                    "ON st.id = a.student_id WHERE a.session_id = %s", (sid2,))}
    assert marked == {"2021104567"}


def test_private_window_with_friends_carne_goes_to_review():
    """Attack 4: a student in the room opens a private window and types a classmate's carne."""
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    friend = student(IPHONE)
    friend.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(friend, t, sid).text),
                                         "carne": "2021104568"})
    sid2 = start(t, sql("SELECT course_id FROM sessions WHERE id = %s", (sid,))[0][0])
    incognito = student(ANDROID_CHROME)
    r = incognito.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(incognito, t, sid2).text),
                                                "carne": "2021104568"})
    assert "Asistencia recibida" in r.text  # not blocked...
    state = t.get(f"/session/{sid2}/state").json()
    labels = [a["label"] for a in state["alerts"]]
    assert main.EVENT_LABELS["cambio_dispositivo"] in labels  # ...but visible
    assert "Chrome" in state["alerts"][0]["detail"]


def test_class_without_list_builds_itself_and_needs_approval():
    t = teacher()
    cid = create_course(t, roster="")
    sid = start(t, cid)
    s = student()
    page = scan(s, t, sid).text
    assert 'name="nombre"' in page
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(page), "carne": "2021200001",
                                        "nombre": "Santiago Herrera"})
    assert "Asistencia recibida" in r.text
    assert sql("SELECT status FROM enrollments")[0][0] == "pendiente"
    t.post(f"/course/{cid}/approve", data={"all": "1"})
    assert sql("SELECT status FROM enrollments")[0][0] == "aprobado"


def test_burst_of_new_links_is_flagged():
    t = teacher()
    sid = start(t, create_course(t, roster=""))
    for i in range(3):
        s = student(ANDROID_CHROME)
        s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                        "carne": f"20213000{i}", "nombre": f"Estudiante {i}"})
    assert sql("SELECT COUNT(*) FROM events WHERE kind = 'rafaga' AND review")[0][0] == 1


def test_link_opened_inside_instagram_is_flagged():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student(INSTAGRAM)
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "2021104569"})
    assert sql("SELECT detail FROM events WHERE kind = 'desde_app'")[0][0] == "Instagram"
    assert sql("SELECT in_app, browser, os FROM attendance")[0] == ("Instagram", "Safari", "iOS")


def test_closed_class_accepts_nothing():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    page = scan(s, t, sid).text
    t.post(f"/session/{sid}/end", follow_redirects=False)
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(page), "carne": "2021104567"})
    assert "Vuelve a escanear" in r.text or "ya terminó" in r.text
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 0


def test_manual_mark_and_reports():
    t = teacher()
    cid = create_course(t, roster=ROSTER)
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "2021104567"})
    juan = sql("SELECT id FROM students WHERE carne = '2021104568'")[0][0]
    assert "marcado a mano" in t.post(f"/session/{sid}/manual", data={"student_id": juan}).json()["message"]
    t.post(f"/session/{sid}/end", follow_redirects=False)

    page = t.get(f"/session/{sid}/report").text
    assert "2 de 3" in page and "Firma del profesor" in page

    rows = list(csv.reader(io.StringIO(t.get(f"/session/{sid}/export.csv").content.decode("utf-8-sig"))))
    header, data = rows[0], rows[1:]
    assert len(header) == len(main.REPORT_COLUMNS)
    by_name = {r[1]: dict(zip(header, r)) for r in data}
    assert by_name["Pérez, Juan"]["Estado"] == "Manual"      # the comma did not shift columns
    assert by_name["Valentina Castro"]["Estado"] == "Ausente"
    assert by_name["María Fernanda Rojas"]["Celular"] == "Celular #1"
    device = sql("SELECT device_id FROM students WHERE carne = '2021104567'")[0][0]
    assert device not in t.get(f"/session/{sid}/export.csv").text  # never the cookie value

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(t.get(f"/session/{sid}/export.xlsx").content))
    assert wb.active.max_row == 4


def test_teachers_cannot_see_each_others_classes():
    a = teacher("a@uni.edu", "Prof. A")
    cid = create_course(a, roster=ROSTER)
    b = teacher("b@uni.edu", "Prof. B")
    r = b.get(f"/course/{cid}", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    sid = start(a, cid)
    assert b.get(f"/session/{sid}/qr", follow_redirects=False).status_code == 303
    assert b.get(f"/session/{sid}/export.csv", follow_redirects=False).status_code == 303


def test_student_typed_name_is_escaped_on_teacher_pages():
    t = teacher()
    cid = create_course(t, roster="")
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "2021400001", "nombre": "<script>alert(1)</script>"})
    page = t.get(f"/course/{cid}").text
    assert "<script>alert(1)</script>" not in page
    assert "&lt;script&gt;" in page


def test_qr_is_teacher_only():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    assert student().get(f"/session/{sid}/qr", follow_redirects=False).status_code == 303


def test_roster_from_excel_and_semicolon_csv():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Código", "Apellidos y nombres"])
    wb.active.append([2021104567, "Rojas María"])
    wb.active.append([2021104568.0, "Pérez Juan"])
    buf = io.BytesIO(); wb.save(buf)
    t = teacher()
    r = t.post("/course/new", data={"name": "Excel"},
               files={"roster_file": ("lista.xlsx", buf.getvalue())}, follow_redirects=False)
    assert r.status_code == 303
    assert sorted(x[0] for x in sql("SELECT carne FROM students")) == ["2021104567", "2021104568"]
    t.post("/course/new", data={"name": "CSV"},
           files={"roster_file": ("lista.csv", "nombre;carné\nAna Gómez;2021104570\n".encode("latin-1"))})
    assert sql("SELECT name FROM students WHERE carne = '2021104570'")[0][0] == "Ana Gómez"


def test_brightspace_export_with_several_groups():
    """The Brightspace grade export: OrgDefinedId, Last Name, First Name, Email, Sections."""
    import json
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["OrgDefinedId", "Last Name", "First Name", "Email", "Sections"])
    ws.append(["170001", "GOMEZ DIAZ", "ANA", "ana@uni.edu", "FUNDAMENTOS MD2026303604"])
    ws.append(["170002", "PEREZ ROA", "LUIS", "luis@uni.edu", "FUNDAMENTOS MD2026303604"])
    ws.append(["170003", "MARIN SOTO", "EVA", "eva@uni.edu", "FUNDAMENTOS MD2026303605"])
    buf = io.BytesIO(); wb.save(buf)
    t = teacher()
    r = t.post("/course/new", data={"name": "Fundamentos"},
               files={"roster_file": ("export.xlsx", buf.getvalue())})
    assert "¿Qué grupos son tuyos?" in r.text and "2 grupos" in r.text and "Grupo 3604" in r.text
    payload = re.search(r'name="payload" value="([^"]+)"', r.text).group(1)
    import html as _html
    payload = _html.unescape(payload)
    assert "@" not in payload  # emails never travel further
    t.post("/course/new/groups", data={"name": "Fundamentos", "payload": payload,
                                       "sections": ["FUNDAMENTOS MD2026303604"]})
    names = sorted(x[0] for x in sql("SELECT name FROM students"))
    assert names == ["Ana Gomez Diaz", "Luis Perez Roa"]



def test_choose_todays_groups_when_class_has_several():
    """One list, several groups: 7:00 with group 3614, 9:00 with group 3615."""
    t = teacher()
    roster = ("OrgDefinedId,Last Name,First Name,Sections\n"
              "180001,ROJAS,ANA,SISTEMA NERVIOSO2026303614\n"
              "180002,DIAZ,LUIS,SISTEMA NERVIOSO2026303614\n"
              "180003,SOTO,EVA,SISTEMA NERVIOSO2026303615\n")
    r = t.post("/course/new", data={"name": "Sistema Nervioso", "roster_text": roster})
    import html as _html
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    r = t.post("/course/new/groups", follow_redirects=False,
               data={"name": "Sistema Nervioso", "payload": payload,
                     "sections": ["SISTEMA NERVIOSO2026303614", "SISTEMA NERVIOSO2026303615"]})
    cid = int(re.search(r"/course/(\d+)", r.headers["location"]).group(1))
    assert sql("SELECT COUNT(*) FROM enrollments")[0][0] == 3

    picker = t.post(f"/course/{cid}/start")
    assert "¿Qué grupos están en el salón ahora?" in picker.text and "Grupo 3615" in picker.text
    r = t.post(f"/course/{cid}/start", follow_redirects=False,
               data={"chosen": "1", "sections": ["SISTEMA NERVIOSO2026303614"]})
    sid = int(re.search(r"/session/(\d+)/live", r.headers["location"]).group(1))

    state = t.get(f"/session/{sid}/state").json()
    assert state["total"] == 2
    assert sorted(a["name"] for a in state["absent"]) == ["Ana Rojas", "Luis Diaz"]

    # Eva, from the other group, shows up anyway: she is marked and flagged in the report.
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "180003"})
    rows = list(csv.reader(io.StringIO(t.get(f"/session/{sid}/export.csv").content.decode("utf-8-sig"))))
    by_name = {r[1]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert set(by_name) == {"Ana Rojas", "Luis Diaz", "Eva Soto"}
    assert "otro grupo" in by_name["Eva Soto"]["Observación"]
    assert by_name["Ana Rojas"]["Grupo"] == "Grupo 3614"


def test_delete_class_needs_its_name_and_keeps_shared_students():
    t = teacher()
    test_class = create_course(t, name="Prueba", roster="Ana Uno\t2021900001\nLuis Dos\t2021900002")
    real = create_course(t, name="Genética", roster="Luis Dos\t2021900002")
    sid = start(t, test_class)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "2021900001"})

    r = t.post(f"/course/{test_class}/delete", data={"confirm_name": "otra cosa"}, follow_redirects=False)
    assert "error=1" in r.headers["location"]
    assert sql("SELECT COUNT(*) FROM courses")[0][0] == 2

    t.post(f"/course/{test_class}/delete", data={"confirm_name": "prueba"})
    assert [r[0] for r in sql("SELECT name FROM courses")] == ["Genética"]
    # Ana belonged only to the test class: gone. Luis is also in Genética: kept.
    assert [r[0] for r in sql("SELECT carne FROM students")] == ["2021900002"]
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 0
    assert t.get(f"/course/{real}").status_code == 200


def test_unlink_phone_lets_the_real_student_link_without_alert():
    """A carne linked to the wrong phone during a test, then unlinked by the professor."""
    t = teacher()
    cid = create_course(t, roster=ROSTER)
    sid = start(t, cid)
    test_phone = student()
    test_phone.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(test_phone, t, sid).text),
                                             "carne": "2021104567"})
    sid_id = sql("SELECT id FROM students WHERE carne = '2021104567'")[0][0]
    t.post(f"/course/{cid}/unlink", data={"student_id": sid_id})
    assert sql("SELECT device_id FROM students WHERE id = %s", (sid_id,))[0][0] is None

    sid2 = start(t, cid)
    real_phone = student(ANDROID_CHROME)
    page = scan(real_phone, t, sid2).text
    assert "Primera vez en este celular" in page
    real_phone.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(page), "carne": "2021104567"})
    assert t.get(f"/session/{sid2}/state").json()["alerts"] == []
    # The test phone is no longer anyone's: it is asked for a carne again.
    assert "Primera vez en este celular" in scan(test_phone, t, sid2).text


def test_professor_corrects_a_mistyped_carne():
    t = teacher()
    cid = create_course(t, roster="")
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "202110456", "nombre": "Juan Perez"})  # missing a digit
    wrong = sql("SELECT id FROM students WHERE carne = '202110456'")[0][0]
    t.post(f"/course/{cid}/student/{wrong}/edit", data={"carne": "2021104567", "name": "Juan Pérez"})
    assert sql("SELECT carne, name FROM students") == [("2021104567", "Juan Pérez")]
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 1
    # The same phone keeps working, now as the right carne.
    sid2 = start(t, cid)
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid2).text)})
    assert "2021104567" in r.text


def test_mistyped_carne_merges_into_the_listed_student():
    t = teacher()
    other = create_course(t, name="Con lista", roster="Ana Gómez\t2021500001")
    cid = create_course(t, name="Sin lista", roster="")
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "202150000", "nombre": "Ana G"})
    wrong = sql("SELECT id FROM students WHERE carne = '202150000'")[0][0]
    t.post(f"/course/{cid}/student/{wrong}/edit", data={"carne": "2021500001", "name": "Ana G"})
    rows = sql("SELECT carne, name, device_id IS NOT NULL FROM students")
    assert rows == [("2021500001", "Ana Gómez", True)]  # one record, the list's name, phone kept
    assert sql("SELECT COUNT(*) FROM enrollments")[0][0] == 2
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 1


def test_professor_rejects_a_flagged_checkin_and_it_stays_on_record():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    friend = student(IPHONE)
    friend.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(friend, t, sid).text),
                                         "carne": "2021104568"})
    sid2 = start(t, sql("SELECT course_id FROM sessions WHERE id = %s", (sid,))[0][0])
    incognito = student(ANDROID_CHROME)
    incognito.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(incognito, t, sid2).text),
                                            "carne": "2021104568"})
    alert = next(a for a in t.get(f"/session/{sid2}/state").json()["alerts"] if a["student"])
    assert alert["decision"] == ""
    assert t.post(f"/session/{sid2}/event/{alert['id']}/decide",
                  data={"decision": "rechazado"}).status_code == 200
    assert t.get(f"/session/{sid2}/state").json()["alerts"][0]["decision"] == "rechazado"
    rows = list(csv.DictReader(io.StringIO(t.get(f"/session/{sid2}/export.csv").text.lstrip("\ufeff"))))
    juan = next(r for r in rows if r["Carné"] == "2021104568")
    assert juan["Estado"] == "Rechazado" and "rechazado por Prof. Rojas" in juan["Observación"]
    assert sql("SELECT COUNT(*) FROM attendance WHERE session_id = %s", (sid2,))[0][0] == 1
    other = teacher("otro@uni.edu")  # another professor cannot decide
    assert other.post(f"/session/{sid2}/event/{alert['id']}/decide",
                      data={"decision": "confirmado"}, follow_redirects=False).status_code in (303, 401)


def test_known_phone_not_on_this_list_is_pending_not_refused():
    """A phone already linked in another class is never sent away from this one."""
    t = teacher()
    other = create_course(t, name="Otra", roster="Ana Gómez\t2021600001")
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, start(t, other)).text),
                                    "carne": "2021600001"})
    sid = start(t, create_course(t, roster=ROSTER))
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text)})
    assert "Asistencia recibida" in r.text
    assert sql("SELECT en.status FROM enrollments en JOIN sessions se ON se.course_id = en.course_id "
               "JOIN students st ON st.id = en.student_id "
               "WHERE se.id = %s AND st.carne = '2021600001'", (sid,))[0][0] == "pendiente"
    labels = [a["label"] for a in t.get(f"/session/{sid}/state").json()["alerts"]]
    assert main.EVENT_LABELS["carne_no_en_lista"] in labels


def test_timing_is_kept_after_class_ends_and_export_is_anonymous():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "2021104567"})
    t.post(f"/session/{sid}/end", follow_redirects=False)
    age, accept, first = sql("SELECT qr_age_s, accept_s, first_link FROM attendance")[0]
    assert 0 <= age <= main.NONCE_VALID_SECONDS and accept >= 0 and first is True
    text = t.get(f"/session/{sid}/timing.csv").text
    rows = list(csv.DictReader(io.StringIO(text)))
    assert len(rows) == 1 and rows[0]["first_link"] == "1"
    assert "2021104567" not in text and "María" not in text
    assert student().get(f"/session/{sid}/timing.csv", follow_redirects=False).status_code == 303


def test_brightspace_hash_prefix_is_ignored():
    """Brightspace writes usernames and ids as "#agomez" and "#170001"."""
    roster = ("OrgDefinedId\tUsername\tLast Name\tFirst Name\tEmail\n"
              "#170001\t#agomez\tGOMEZ DIAZ\tANA\tagomez@universidad.edu.co\n")
    t = teacher()
    sid = start(t, create_course(t, roster=roster))
    assert sql("SELECT carne, username FROM students") == [("170001", "AGOMEZ")]
    s = student()
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text), "carne": "agomez"})
    assert "Asistencia recibida" in r.text and "Ana Gomez Diaz" in r.text


def test_nothing_is_stored_without_accepting_the_data_policy():
    t = teacher()
    sid = start(t, create_course(t, roster=""))
    s = student()
    tk = ticket_of(scan(s, t, sid).text)
    r = s.post("/checkin/accept", data={"ticket": tk, "carne": "2021700001", "nombre": "Ana Ruiz"})
    assert "acepta la política" in r.text and 'name="acepto"' in r.text
    assert sql("SELECT COUNT(*) FROM students")[0][0] == 0
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 0
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": tk, "carne": "2021700001",
                                        "nombre": "Ana Ruiz"})
    assert "Asistencia recibida" in r.text
    assert sql("SELECT consent_at IS NOT NULL FROM students")[0][0] is True
    assert "Ley 1581" in s.get("/privacidad").text


def test_class_without_list_adds_a_known_phone_as_pending():
    """Someone who linked a phone in another class just taps Aceptar here."""
    t = teacher()
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, start(t, create_course(t, roster=ROSTER))).text),
                                    "carne": "2021104567"})
    cid = create_course(t, name="Congreso", roster="")
    r = s.post("/checkin/accept", data={"ticket": ticket_of(scan(s, t, start(t, cid)).text)})
    assert "Asistencia recibida" in r.text and "María Fernanda Rojas" in r.text
    assert sql("SELECT en.status FROM enrollments en WHERE en.course_id = %s", (cid,)) == [("pendiente",)]


def test_a_typed_name_cannot_become_an_excel_formula():
    t = teacher()
    sid = start(t, create_course(t, roster=""))
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "2021800001", "nombre": '=HYPERLINK("http://x.co","ver")'})
    rows = list(csv.reader(io.StringIO(t.get(f"/session/{sid}/export.csv").content.decode("utf-8-sig"))))
    assert rows[1][1].startswith("'=")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(t.get(f"/session/{sid}/export.xlsx").content))
    assert wb.active.cell(row=2, column=2).data_type == "s"


def test_banner_export_with_groups():
    """A Banner class list: SPRIDEN_ID, SPRIDEN_LAST_NAME, SSBSECT_CRN (fictional rows)."""
    import html as _html
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SPRIDEN_ID", "SPRIDEN_LAST_NAME", "SSBSECT_CRN"])
    ws.append([200100001, "ANA GOMEZ DIAZ", 3614])
    ws.append([200100002, "LUIS PEREZ ROA", 3614])
    ws.append([200100003, "EVA MARIN SOTO", 3615])
    buf = io.BytesIO(); wb.save(buf)
    t = teacher()
    r = t.post("/course/new", data={"name": "Sistema Nervioso"},
               files={"roster_file": ("Listado.xlsx", buf.getvalue())})
    assert "Grupo 3614" in r.text and "Grupo 3615" in r.text
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    r = t.post("/course/new/groups", follow_redirects=False,
               data={"name": "Sistema Nervioso", "payload": payload, "sections": ["3614", "3615"]})
    assert sorted(x[0] for x in sql("SELECT name FROM students")) == ["Ana Gomez Diaz", "Eva Marin Soto",
                                                                     "Luis Perez Roa"]
    assert sorted(x[0] for x in sql("SELECT section FROM enrollments")) == ["3614", "3614", "3615"]
    cid = int(re.search(r"/course/(\d+)", r.headers["location"]).group(1))
    sid = int(re.search(r"/session/(\d+)/live", t.post(f"/course/{cid}/start", follow_redirects=False,
              data={"chosen": "1", "sections": ["3614"]}).headers["location"]).group(1))
    s = student()
    r = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                        "carne": "200100001"})
    assert "Asistencia recibida" in r.text and "Ana Gomez Diaz" in r.text


def test_attendance_without_list_is_validated_when_the_list_arrives():
    """Real case, 21 Oct 2026: attendance taken with no list, the Brightspace list uploaded after. Fictional names."""
    import html as _html
    t = teacher()
    cid = create_course(t, name="Fundamentos", roster="")
    sid = start(t, cid)
    typed = [("200100101", "Camila Reyes Duarte", IPHONE),        # her code: matches by name
             ("SMORA", "Sofia Mora Vidal", ANDROID_CHROME),  # her username: exact
             ("200100103", "María Lucía Navarro", INSTAGRAM),  # only part of her name
             ("200999999", "Persona Ajena", IPHONE)]              # not on the list
    phones = {}
    for carne, name, ua in typed:
        s = student(ua)
        s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                        "carne": carne, "nombre": name})
        phones[carne] = s
    t.post(f"/session/{sid}/end", follow_redirects=False)

    roster = ("OrgDefinedId,Last Name,First Name,Email,Sections\n"
              "170001,REYES DUARTE,CAMILA,creyes@universidad.edu.co,FUNDAMENTO MD2026303604\n"
              "170002,MORA VIDAL,SOFIA,smora@universidad.edu.co,FUNDAMENTO MD2026303604\n"
              "170003,NAVARRO PEÑA,MARIA ANDREA,manavarro@universidad.edu.co,FUNDAMENTO MD2026303604\n"
              "170004,OTRO,ALUMNO,otro@universidad.edu.co,FUNDAMENTO MD2026303605\n")
    r = t.post(f"/course/{cid}/roster", data={"roster_text": roster})
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    r = t.post(f"/course/{cid}/roster/groups", follow_redirects=False,
               data={"payload": payload, "sections": ["FUNDAMENTO MD2026303604"]})
    assert "/validar" in r.headers["location"]

    page = t.get(f"/course/{cid}/validar").text
    assert "Mismo nombre" in page and "Mismo código o usuario" in page
    assert "Nombre parecido" in page and "No está en la lista" in page
    pairs = re.findall(r"name='pair' value='(\d+:\d+)' (checked)?", page)
    assert sum(1 for _, c in pairs if c) == 2     # the partial match is not pre-checked
    t.post(f"/course/{cid}/validar", data={"pair": [p for p, c in pairs if c]})

    gil = sql("SELECT id, carne, name, username, device_id IS NOT NULL FROM students "
              "WHERE username = 'CREYES'")[0]
    assert gil[1:] == ("200100101", "Camila Reyes Duarte", "CREYES", True)  # her code kept, list's name
    assert sql("SELECT COUNT(*) FROM attendance WHERE student_id = %s", (gil[0],))[0][0] == 1
    assert sql("SELECT status FROM enrollments WHERE student_id = %s AND course_id = %s",
               (gil[0], cid))[0][0] == "aprobado"
    assert sql("SELECT carne FROM students WHERE username = 'SMORA'")[0][0] == "170002"
    assert sql("SELECT COUNT(*) FROM enrollments WHERE course_id = %s AND status = 'pendiente'",
               (cid,))[0][0] == 2
    # Next class her phone is recognized as the listed student, one tap.
    sid2 = start(t, cid)
    r = phones["200100101"].post("/checkin/accept",
                                 data={"ticket": ticket_of(scan(phones["200100101"], t, sid2).text)})
    assert "Camila Reyes Duarte" in r.text
    # Uploading the same list again creates no duplicates.
    before = sql("SELECT COUNT(*) FROM students")[0][0]
    r = t.post(f"/course/{cid}/roster", data={"roster_text": roster})
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    t.post(f"/course/{cid}/roster/groups", data={"payload": payload, "sections": ["FUNDAMENTO MD2026303604"]})
    assert sql("SELECT COUNT(*) FROM students")[0][0] == before


def test_fix_groups_after_uploading_too_many():
    """Real case, 21 Oct 2026: every group was uploaded, but only one was in the room."""
    import html as _html
    t = teacher()
    cid = create_course(t, name="Prueba", roster="")
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "200100101", "nombre": "Camila Reyes Duarte"})
    t.post(f"/session/{sid}/end", follow_redirects=False)
    roster = ("OrgDefinedId,Last Name,First Name,Email,Sections\n"
              "170001,REYES DUARTE,CAMILA,creyes@universidad.edu.co,FUNDAMENTO MD2026303611\n"
              "170002,MORA VIDAL,SOFIA,smora@universidad.edu.co,FUNDAMENTO MD2026303611\n"
              "170003,PAZ LEON,TOMAS,tpaz@universidad.edu.co,FUNDAMENTO MD2026303604\n"
              "170004,RUIZ SOL,ANA,aruiz@universidad.edu.co,FUNDAMENTO MD2026303605\n")
    r = t.post(f"/course/{cid}/roster", data={"roster_text": roster})
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    t.post(f"/course/{cid}/roster/groups", data={"payload": payload, "sections": [
        "FUNDAMENTO MD2026303611", "FUNDAMENTO MD2026303604", "FUNDAMENTO MD2026303605"]})
    pairs = re.findall(r"name='pair' value='(\d+:\d+)' checked", t.get(f"/course/{cid}/validar").text)
    t.post(f"/course/{cid}/validar", data={"pair": pairs})
    assert "1 de 4" in t.get(f"/session/{sid}/report").text          # the mistake

    page = t.get(f"/session/{sid}/report").text
    assert "¿Qué grupos estaban en esta sesión?" in page
    t.post(f"/session/{sid}/groups", data={"sections": ["FUNDAMENTO MD2026303611"]})
    assert "1 de 2" in t.get(f"/session/{sid}/report").text          # fixed for that session

    t.post(f"/course/{cid}/groups/remove", data={"section": "FUNDAMENTO MD2026303604"})
    t.post(f"/course/{cid}/groups/remove", data={"section": "FUNDAMENTO MD2026303605"})
    assert sql("SELECT COUNT(*) FROM enrollments WHERE course_id = %s", (cid,))[0][0] == 2
    assert sql("SELECT COUNT(*) FROM attendance WHERE session_id = %s", (sid,))[0][0] == 1
    other = teacher("otro@uni.edu")
    assert other.post(f"/session/{sid}/groups", data={"sections": []},
                      follow_redirects=False).status_code == 303
    assert sql("SELECT sections FROM sessions WHERE id = %s", (sid,))[0][0] == ["FUNDAMENTO MD2026303611"]


def test_rename_class_sort_by_surname_and_burst_names_students():
    t = teacher()
    roster = ("OrgDefinedId,Last Name,First Name\n"
              "170001,ZAPATA,ANA\n170002,ARIAS,LUIS\n170003,MEJIA,EVA\n")
    cid = create_course(t, name="Prueba", roster=roster)
    t.post(f"/course/{cid}/rename", data={"name": "Fundamentos"})
    assert sql("SELECT name FROM courses WHERE id = %s", (cid,))[0][0] == "Fundamentos"
    sid = start(t, cid)
    for carne in ("170001", "170002", "170003"):
        s = student()
        s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                        "carne": carne})
    page = t.get(f"/session/{sid}/report").text
    assert "data-apellido='arias luis arias'" in page and 'Apellido, Z a A' in page and "filterReport('Ausente'" in page
    assert "Ana Zapata" in sql("SELECT detail FROM events WHERE kind = 'rafaga'")[0][0]
    assert "Varios estudiantes" in page


def test_second_list_with_other_codes_updates_instead_of_duplicating():
    """Brightspace list first (6-digit ids), Banner list later (9-digit codes)."""
    t = teacher()
    cid = create_course(t, roster="OrgDefinedId,Last Name,First Name,Email\n"
                                   "430465,VERGARA ROMERO,CARLOS,cvergara@u.edu\n"
                                   "431830,MORA VIDAL,SOFIA,smora@u.edu\n")
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "smora"})
    t.post(f"/course/{cid}/roster", data={"roster_text": "SPRIDEN_ID\tSPRIDEN_LAST_NAME\n"
                                          "200300001\tVERGARA ROMERO CARLOS\n"
                                          "200300002\tMORA VIDAL SOFIA\n"})
    assert sorted(r[0] for r in sql("SELECT carne FROM students")) == ["200300001", "200300002"]
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 1
    assert sql("SELECT username FROM students WHERE carne = '200300002'")[0][0] == "SMORA"


def test_groups_are_chosen_on_the_class_page_and_removal_is_tucked_away():
    t = teacher()
    roster = ("OrgDefinedId,Last Name,First Name,Sections\n"
              "1,ROJAS,ANA,MAT2026303614\n2,DIAZ,LUIS,MAT2026303615\n")
    r = t.post("/course/new", data={"name": "Mat", "roster_text": roster})
    import html as _html
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    r = t.post("/course/new/groups", follow_redirects=False,
               data={"name": "Mat", "payload": payload, "sections": ["MAT2026303614", "MAT2026303615"]})
    cid = int(re.search(r"/course/(\d+)", r.headers["location"]).group(1))
    page = t.get(f"/course/{cid}").text
    start_form = page.split('action="/course/%d/start"' % cid)[1].split("</form>")[0]
    assert "name='sections'" in start_form and "name='chosen'" in start_form
    assert "Quitar un grupo de esta clase" in page
    # nothing checked: attendance starts with every group
    r = t.post(f"/course/{cid}/start", data={"chosen": "1"}, follow_redirects=False)
    sid = int(re.search(r"/session/(\d+)/live", r.headers["location"]).group(1))
    assert t.get(f"/session/{sid}/state").json()["total"] == 2


def test_history_has_its_own_page_and_cards_link_both():
    t = teacher()
    cid = create_course(t, name="Gen", roster=ROSTER)
    sid = start(t, cid)
    t.post(f"/session/{sid}/end", follow_redirects=False)
    home = t.get("/").text
    assert f"/course/{cid}/historial" in home and "Ver lista" in home
    page = t.get(f"/course/{cid}/historial").text
    assert "Historial" in page and f"/session/{sid}/report" in page and "Tomar asistencia</button>" not in page


def test_delete_one_test_session_and_the_data_it_created():
    t = teacher()
    cid = create_course(t, name="Fundamentos", roster=ROSTER)
    real = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, real).text),
                                    "carne": "2021104567"})
    t.post(f"/session/{real}/end", follow_redirects=False)
    test = start(t, cid)
    me = student(ANDROID_CHROME)
    me.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(me, t, test).text),
                                     "carne": "1045678901", "nombre": "Probador"})
    t.post(f"/session/{test}/end", follow_redirects=False)
    assert f"/session/{test}/delete" in t.get(f"/course/{cid}/historial").text
    r = t.post(f"/session/{test}/delete", data={"confirm_id": "999"}, follow_redirects=False)
    assert "error=1" in r.headers["location"]
    t.post(f"/session/{test}/delete", data={"confirm_id": str(test)})
    assert sql("SELECT COUNT(*) FROM sessions WHERE id = %s", (test,))[0][0] == 0
    assert sql("SELECT COUNT(*) FROM students WHERE carne = '1045678901'")[0][0] == 0
    assert sql("SELECT COUNT(*) FROM attendance WHERE session_id = %s", (real,))[0][0] == 1
    assert sql("SELECT COUNT(*) FROM students WHERE carne = '2021104567'")[0][0] == 1
    other = teacher("otro@uni.edu")
    assert other.post(f"/session/{real}/delete", data={"confirm_id": str(real)},
                      follow_redirects=False).status_code == 303
    assert sql("SELECT COUNT(*) FROM sessions WHERE id = %s", (real,))[0][0] == 1


def test_removing_a_student_from_their_only_class_deletes_their_data():
    t = teacher()
    cid = create_course(t, name="Prueba", roster="")
    sid = start(t, cid)
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "1045678901", "nombre": "Probador"})
    st = sql("SELECT id FROM students WHERE carne = '1045678901'")[0][0]
    t.post(f"/course/{cid}/remove", data={"student_id": st})
    assert sql("SELECT COUNT(*) FROM students WHERE id = %s", (st,))[0][0] == 1   # has attendance
    t.post(f"/session/{sid}/end", follow_redirects=False)
    t.post(f"/session/{sid}/delete", data={"confirm_id": str(sid)})
    assert sql("SELECT COUNT(*) FROM students WHERE id = %s", (st,))[0][0] == 0


def _all_at_once(fn, n):
    """Run fn(i) for i in range(n) in n threads released at the same instant."""
    import threading
    barrier = threading.Barrier(n)
    out = [None] * n

    def worker(i):
        barrier.wait()
        out[i] = fn(i)

    threads = [threading.Thread(target=worker, args=(i,)) for i in range(n)]
    for th in threads:
        th.start()
    for th in threads:
        th.join()
    return out


def test_a_whole_class_tapping_at_once_loses_and_doubles_nobody():
    """A class of 100 taps Aceptar within seconds. Simulated here with 30 at the same instant."""
    n = 30
    roster = "\n".join(f"Estudiante Prueba {i}\t{300000000 + i}" for i in range(n))
    t = teacher()
    sid = start(t, create_course(t, roster=roster))
    nonce = fresh_nonce(t, sid)
    phones = [student() for _ in range(n)]
    tickets = [ticket_of(p.get(f"/checkin?s={sid}&n={nonce}").text) for p in phones]
    pages = _all_at_once(lambda i: phones[i].post("/checkin/accept", data={
        "acepto": "1", "ticket": tickets[i], "carne": str(300000000 + i)}).text, n)
    assert all("Asistencia recibida" in p for p in pages)
    assert sql("SELECT COUNT(*), COUNT(DISTINCT student_id) FROM attendance "
               "WHERE session_id = %s", (sid,))[0] == (n, n)


def test_double_tap_on_aceptar_creates_one_attendance():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    tk = ticket_of(scan(s, t, sid).text)
    pages = _all_at_once(lambda i: s.post("/checkin/accept", data={
        "acepto": "1", "ticket": tk, "carne": "2021104567"}).text, 2)
    assert sum("Asistencia recibida" in p for p in pages) == 1
    assert sql("SELECT COUNT(*) FROM attendance")[0][0] == 1


def test_projected_screen_hides_names_and_report_shows_alerts_first():
    """The live screen is projected: names stay hidden until the professor opens them."""
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    page = t.get(f"/session/{sid}/live").text
    assert re.search(r'<div id="details"[^>]*hidden>', page)
    assert "Mostrar ausentes y alertas" in page
    report = t.get(f"/session/{sid}/report").text
    assert report.index('id="por-revisar"') < report.index('id="report"')


def test_list_with_only_codes_asks_for_the_code():
    """Real case, 26 Oct 2026: a Banner list has codes, not usernames, so the form asks for the code."""
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    page = scan(student(), t, sid).text
    assert "Código de estudiante" in page and "Carné o usuario" not in page
    roster = ("OrgDefinedId,Last Name,First Name,Email\n"
              "170001,GOMEZ DIAZ,ANA,agomez@universidad.edu.co\n")
    sid2 = start(t, create_course(t, name="Con usuarios", roster=roster))
    assert "Carné o usuario" in scan(student(), t, sid2).text


def test_expired_scan_names_a_known_phone_and_the_professor_can_mark_after_class():
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "2021104567"})
    t.post(f"/session/{sid}/end")
    sid2 = start(t, sql("SELECT course_id FROM sessions WHERE id = %s", (sid,))[0][0])
    n = fresh_nonce(t, sid2)
    sql("UPDATE nonces SET expires_at = now() - interval '1 second' WHERE nonce = %s", (n,))
    s.get(f"/checkin?s={sid2}&n={n}")
    t.post(f"/session/{sid2}/end")
    report = t.get(f"/session/{sid2}/report").text
    assert "María Fernanda Rojas" in report and "no marcó" in report
    student_id = sql("SELECT id FROM students WHERE carne = '2021104567'")[0][0]
    r = t.post(f"/session/{sid2}/manual", data={"student_id": student_id}).json()
    assert "marcado a mano" in r["message"]
    assert "después marcó" in t.get(f"/session/{sid2}/report").text


def test_report_lists_alerts_about_one_student_before_the_general_notice():
    t = teacher()
    sid = start(t, create_course(t, roster=""))
    for i in range(3):
        s = student(ANDROID_CHROME)
        s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                        "carne": f"20213000{i}", "nombre": f"Estudiante {i}"})
    insta = student(INSTAGRAM)
    insta.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(insta, t, sid).text),
                                        "carne": "202130009", "nombre": "Estudiante Nueve"})
    report = t.get(f"/session/{sid}/report").text
    assert report.index("Abrió el enlace dentro de una app") < report.index("Avisos generales")


def test_a_username_approved_by_hand_can_still_be_joined_to_its_list_row():
    """Real case, 26 Oct 2026: "Aprobar" instead of "Validar" left each student twice."""
    import html as _html
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SPRIDEN_ID", "SPRIDEN_LAST_NAME", "SSBSECT_CRN"])
    ws.append([200100001, "ANA GOMEZ DIAZ", 3614])
    ws.append([200100002, "LUIS PEREZ ROA", 3614])
    ws.append([200100003, "EVA MARIN SOTO", 3615])
    buf = io.BytesIO(); wb.save(buf)
    t = teacher()
    r = t.post("/course/new", data={"name": "Sistema Nervioso"},
               files={"roster_file": ("Listado.xlsx", buf.getvalue())})
    payload = _html.unescape(re.search(r'name="payload" value="([^"]+)"', r.text).group(1))
    r = t.post("/course/new/groups", follow_redirects=False,
               data={"name": "Sistema Nervioso", "payload": payload, "sections": ["3614", "3615"]})
    cid = int(re.search(r"/course/(\d+)", r.headers["location"]).group(1))
    sid = int(re.search(r"/session/(\d+)/live", t.post(f"/course/{cid}/start", follow_redirects=False,
              data={"chosen": "1", "sections": ["3614"]}).headers["location"]).group(1))
    s = student()
    page = s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                           "carne": "agomez"}).text
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "agomez", "nombre": "Ana Gomez"})
    wrong = sql("SELECT id FROM students WHERE carne = 'AGOMEZ'")[0][0]
    t.post(f"/course/{cid}/approve", data={"student_id": wrong})
    assert sql("SELECT COUNT(*) FROM enrollments WHERE course_id = %s", (cid,))[0][0] == 4
    assert "1 estudiantes aprobados parecen estar" in t.get(f"/course/{cid}").text
    page = t.get(f"/course/{cid}/validar").text
    right = sql("SELECT id FROM students WHERE carne = '200100001'")[0][0]
    assert f"{wrong}:{right}" in page and "Luis" not in page and "Eva" not in page
    t.post(f"/course/{cid}/validar", data={"pair": [f"{wrong}:{right}"]})
    assert sql("SELECT COUNT(*) FROM enrollments WHERE course_id = %s", (cid,))[0][0] == 3
    assert sql("SELECT carne, username FROM students WHERE id = %s", (right,)) == [("200100001", "AGOMEZ")]
    assert sql("SELECT COUNT(*) FROM attendance WHERE student_id = %s", (right,))[0][0] == 1
    report = t.get(f"/session/{sid}/report").text
    assert "<b>0</b>por revisar" in report and "Nada que revisar en esta clase" in report
    assert "Ya resueltas al validar con la lista (1)" in report


def test_old_not_on_list_alerts_get_their_name_back():
    """Alerts saved before October 27 had no student; the upgrade names them by the code typed."""
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    student_id = sql("SELECT id FROM students WHERE carne = '2021104567'")[0][0]
    sql("INSERT INTO events (session_id, student_id, kind, detail, review) "
        "VALUES (%s, NULL, 'carne_no_en_lista', '2021104567 · Safari · iOS', TRUE)", (sid,))
    main.init_db()
    assert sql("SELECT student_id FROM events WHERE kind = 'carne_no_en_lista'") == [(student_id,)]


def test_professor_marks_a_late_student_with_a_note_and_timing_has_seconds():
    """A student who arrives mid-class is marked by hand with a note, without reopening the QR."""
    t = teacher()
    sid = start(t, create_course(t, roster=ROSTER))
    s = student()
    s.post("/checkin/accept", data={"acepto": "1", "ticket": ticket_of(scan(s, t, sid).text),
                                    "carne": "2021104567"})
    t.post(f"/session/{sid}/end")
    report = t.get(f"/session/{sid}/report").text
    assert report.count("Marcar presente</button>") >= 2   # the two absent students
    juan = sql("SELECT id FROM students WHERE carne = '2021104568'")[0][0]
    t.post(f"/session/{sid}/manual", data={"student_id": juan, "note": "Llegó a las 9:10"})
    report = t.get(f"/session/{sid}/report").text
    assert "Nota: Llegó a las 9:10" in report
    rows = list(csv.DictReader(io.StringIO(t.get(f"/session/{sid}/timing.csv").text)))
    assert len(rows) == 1 and float(rows[0]["seconds_since_start"]) >= 0
