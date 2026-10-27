"""
UniPresence - Privacy-preserving physical presence verification for large classrooms
Version 2: rebuilt from what the pilot of 22 September 2026 showed

WHAT CHANGED, AND WHY

Version 1 ran in a real genetics lab: 22 students, the professor operating it
alone. 15 checked in, 7 could not, and the stored data was lost when Render's
free instance went to sleep. Every change below answers one of those facts, or
one thing the professor had to do by hand.

1. Data lives in Postgres (Neon), not in a SQLite file inside the server.
   A restart, a redeploy or a sleeping instance no longer erases attendance.

2. One QR does everything. Version 1 had a registration QR and an attendance
   QR. The 7 failures came from that split: students registered inside
   WhatsApp, then scanned with the camera, which opens Safari, and the system
   no longer recognised them. Now the attendance QR is the only door. A phone
   the system knows sees its class and taps Aceptar. A phone it does not know
   types the carne once, right there, and is marked in the same step.

3. The professor prepares nothing. Each professor has an account and their
   own classes. A class list can be uploaded or pasted once and is kept; a
   class can also start with no list at all, and the list builds itself from
   the students who check in (they wait as "pendiente" until approved).
   The printed slips, the Python script and the secret file on Render are gone.

4. Fraud is made visible instead of forbidden with secrets. Without private
   codes, a classmate can type someone else's carne. The system cannot stop
   that without collecting data that identifies the phone, so it records what
   is unusual and shows it to the professor under "Por revisar": a carne that
   moves to another browser during class, several new phones linked in a few
   seconds from the same browser and system, a link opened inside WhatsApp or
   Instagram. These are leads, not proof. The professor decides.

5. Every check-in records the browser, the operating system and whether the
   page was opened inside an app. Only those three words are kept, never the
   full User-Agent string, which is detailed enough to help identify a phone.

6. The student gets a receipt: a full-screen green confirmation with a live
   clock and a short receipt code. A screenshot of an old confirmation does
   not show a running clock, and the code can be looked up in the report.

7. Reports: a printable page (saved as PDF from the browser), Excel and CSV.
   The CSV of version 1 was built by joining text with commas, which broke on
   any name that contained one. The reports show "Celular #1", never the real
   device identifier: that value is the student's cookie, and anyone holding
   it could pass as that student.

HOW TO RUN:
    DATABASE_URL="postgresql://..." uvicorn main:app --reload

Optional: TEACHER_SIGNUP_CODE="..." makes new teacher accounts require that
code, so a stranger who finds the site cannot open an account.

The app refuses to start without DATABASE_URL. Failing closed is deliberate:
a missing database must not silently mean "keep everything in memory".
"""

import csv
import html
import io
import json
import os
import re
import secrets
import unicodedata
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import bcrypt
import psycopg
import qrcode
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import (HTMLResponse, JSONResponse, RedirectResponse,
                               Response, StreamingResponse)
from psycopg.rows import dict_row
from psycopg_pool import ConnectionPool

# ---------------------------------------------------------------
# CONFIGURATION
# ---------------------------------------------------------------

DATABASE_URL = os.environ.get("DATABASE_URL", "")

# Version 2 keeps its tables in their own Postgres schema. If the version 1
# tables were already created in the same Neon database, they stay untouched
# and the two never collide.
SCHEMA = "unipresence_v2"

LOCAL_TZ = ZoneInfo("America/Bogota")

# The projected QR changes every 10 seconds, but each code stays valid for 20.
# The extra margin is for the student at the back of a large hall whose camera
# takes a moment to focus. It is checked at the moment of the scan.
QR_ROTATE_SECONDS = 10
NONCE_VALID_SECONDS = 20

# After a valid scan, the student has this long to tap Aceptar or type the
# carne. Typing is slow; version 1 learned on day 1 that a code expiring while
# someone types is a dead end.
TICKET_SECONDS = 180

TEACHER_SESSION_HOURS = 12
SIGNUP_CODE = os.environ.get("TEACHER_SIGNUP_CODE", "")
DEFAULT_LATE_MINUTES = 15

# "Several new phones linked in a few seconds from the same browser and
# system" is one of the review signals. These are its two thresholds.
BURST_WINDOW_SECONDS = 60
BURST_COUNT = 3

# Same alphabet as generate_codes.py: no 0/O, no 1/I/L.
RECEIPT_ALPHABET = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"

app = FastAPI(title="UniPresence")


# ---------------------------------------------------------------
# DATABASE
# ---------------------------------------------------------------

@contextmanager
def db():
    """
    One request, one transaction.

    Neon's connection pooler hands a different server connection to each
    transaction, so a plain `SET search_path` could be lost between
    statements. `SET LOCAL` lives exactly as long as this transaction, which
    is why nothing inside a request commits halfway through.

    prepare_threshold=None turns off server-side prepared statements, which
    a transaction pooler cannot keep between transactions.
    """
    with get_pool().connection() as conn:
        try:
            conn.execute(f"SET LOCAL search_path TO {SCHEMA}, public")
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise


_pool = None


def get_pool() -> ConnectionPool:
    """
    A few connections kept open and reused.

    Opening a connection to Neon means a network round trip and an encrypted
    handshake, a noticeable fraction of a second. The live screen asks the
    server something every 3 seconds and a new QR every 10, so paying that on
    every request made the whole site feel slow. The pool pays it once.

    Neon's free plan suspends the database after 5 minutes without use, which
    drops open connections. `check` tests each connection before handing it
    out and replaces a dead one, and `max_idle` closes connections idle for
    4 minutes, before Neon would.
    """
    global _pool
    if _pool is None:
        _pool = ConnectionPool(
            DATABASE_URL, min_size=1, max_size=5, max_idle=240, open=True,
            check=ConnectionPool.check_connection,
            kwargs={"row_factory": dict_row, "prepare_threshold": None},
        )
    return _pool


SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS teachers (
    id            BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    email         TEXT NOT NULL UNIQUE,
    name          TEXT NOT NULL,
    institution   TEXT NOT NULL DEFAULT '',
    password_hash TEXT NOT NULL,
    created_at    TIMESTAMPTZ NOT NULL DEFAULT now()
);

-- Kept in the database, not in memory, so a restart no longer logs the
-- professor out in the middle of a class.
CREATE TABLE IF NOT EXISTS teacher_sessions (
    token      TEXT PRIMARY KEY,
    teacher_id BIGINT NOT NULL REFERENCES teachers(id) ON DELETE CASCADE,
    expires_at TIMESTAMPTZ NOT NULL
);

CREATE TABLE IF NOT EXISTS courses (
    id           BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    teacher_id   BIGINT NOT NULL REFERENCES teachers(id) ON DELETE CASCADE,
    name         TEXT NOT NULL,
    has_roster   BOOLEAN NOT NULL DEFAULT FALSE,
    created_at   TIMESTAMPTZ NOT NULL DEFAULT now(),
    last_used_at TIMESTAMPTZ
);

-- One row per person, identified by carne, shared by every class. A student
-- in six classes links a phone once, not six times.
-- device_id stays NULL until the student links a browser. Postgres allows
-- several NULLs in a UNIQUE column.
CREATE TABLE IF NOT EXISTS students (
    id            BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    carne         TEXT NOT NULL UNIQUE,
    name          TEXT NOT NULL,
    device_id     TEXT UNIQUE,
    device_no     INTEGER NOT NULL DEFAULT 0,
    registered_at TIMESTAMPTZ
);

-- Who belongs to which class. 'pendiente' is a student who joined a class
-- with no list by checking in, and is waiting for the professor's approval.
CREATE TABLE IF NOT EXISTS enrollments (
    course_id  BIGINT NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    student_id BIGINT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    status     TEXT NOT NULL CHECK (status IN ('aprobado', 'pendiente')),
    section    TEXT NOT NULL DEFAULT '',
    created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    PRIMARY KEY (course_id, student_id)
);

CREATE TABLE IF NOT EXISTS sessions (
    id         BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    course_id  BIGINT NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
    started_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    ended_at   TIMESTAMPTZ,
    active     BOOLEAN NOT NULL DEFAULT TRUE,
    sections   TEXT[]
);

CREATE TABLE IF NOT EXISTS nonces (
    nonce      TEXT PRIMARY KEY,
    session_id BIGINT NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
    expires_at TIMESTAMPTZ NOT NULL
);

-- A valid scan becomes a one-use ticket. The nonce is checked when the
-- student scans; the ticket carries that result to the Aceptar button.
CREATE TABLE IF NOT EXISTS scan_tickets (
    token      TEXT PRIMARY KEY,
    session_id BIGINT NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
    scanned_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    used_at    TIMESTAMPTZ
);

CREATE TABLE IF NOT EXISTS attendance (
    id         BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    session_id BIGINT NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
    student_id BIGINT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    marked_at  TIMESTAMPTZ NOT NULL DEFAULT now(),
    method     TEXT NOT NULL,
    receipt    TEXT NOT NULL UNIQUE,
    browser    TEXT,
    os         TEXT,
    in_app     TEXT,
    device_no  INTEGER,
    UNIQUE (session_id, student_id)
);

-- Every time a carne moves to a different browser or phone. Re-linking is
-- allowed so that a present student is never locked out, as 7 were in the
-- pilot; writing down every move is what keeps that visible.
CREATE TABLE IF NOT EXISTS device_changes (
    id         BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    student_id BIGINT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    old_device TEXT,
    new_device TEXT NOT NULL,
    session_id BIGINT REFERENCES sessions(id) ON DELETE SET NULL,
    changed_at TIMESTAMPTZ NOT NULL DEFAULT now()
);

-- Everything worth knowing that is not an attendance row: links, moves,
-- blocked attempts, expired scans. `review` marks what goes under
-- "Por revisar".
CREATE TABLE IF NOT EXISTS events (
    id         BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    session_id BIGINT REFERENCES sessions(id) ON DELETE CASCADE,
    student_id BIGINT REFERENCES students(id) ON DELETE SET NULL,
    kind       TEXT NOT NULL,
    detail     TEXT NOT NULL DEFAULT '',
    review     BOOLEAN NOT NULL DEFAULT FALSE,
    browser    TEXT,
    os         TEXT,
    in_app     TEXT,
    created_at TIMESTAMPTZ NOT NULL DEFAULT now()
);
"""


# Columns added after the first deploy of version 2. ADD COLUMN IF NOT EXISTS
# upgrades a database created earlier without touching its data.
UPGRADE_SQL = """
ALTER TABLE enrollments ADD COLUMN IF NOT EXISTS section TEXT NOT NULL DEFAULT '';
ALTER TABLE sessions ADD COLUMN IF NOT EXISTS sections TEXT[];
-- Brightspace exports an internal number (OrgDefinedId) that students do not
-- know. Their institutional username, the part of the email before the @,
-- they do know. It is kept as a second way to find the same student.
ALTER TABLE students ADD COLUMN IF NOT EXISTS username TEXT;
CREATE UNIQUE INDEX IF NOT EXISTS students_username_key ON students (username);
-- The professor's decision on each item under "Por revisar". The event is
-- never deleted: the decision is written next to it.
ALTER TABLE events ADD COLUMN IF NOT EXISTS decision TEXT;
ALTER TABLE events ADD COLUMN IF NOT EXISTS decided_at TIMESTAMPTZ;
ALTER TABLE events ADD COLUMN IF NOT EXISTS decided_by TEXT;
-- Timing of each QR check-in, for the relay analysis: how old the QR was
-- when it was scanned, and how long the student took to tap Aceptar.
-- Tickets are deleted when the class ends, so the numbers are copied to the
-- attendance row, which is kept.
ALTER TABLE scan_tickets ADD COLUMN IF NOT EXISTS qr_age_s REAL;
ALTER TABLE attendance ADD COLUMN IF NOT EXISTS qr_age_s REAL;
ALTER TABLE attendance ADD COLUMN IF NOT EXISTS accept_s REAL;
ALTER TABLE attendance ADD COLUMN IF NOT EXISTS first_link BOOLEAN;
-- When the person accepted the data policy (Law 1581 of 2012 asks for prior,
-- informed authorization). Asked once, the first time a phone is linked.
ALTER TABLE students ADD COLUMN IF NOT EXISTS consent_at TIMESTAMPTZ;
-- The surname when the list has it in its own column, so a report can be
-- sorted by last name the way professors read a class list.
ALTER TABLE students ADD COLUMN IF NOT EXISTS last_name TEXT;
"""


def init_db():
    with db() as conn:
        conn.execute(f"CREATE SCHEMA IF NOT EXISTS {SCHEMA}")
        conn.execute(SCHEMA_SQL)
        conn.execute(UPGRADE_SQL)


@app.on_event("startup")
def on_startup():
    if not DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL is not set. UniPresence v2 stores everything in "
            "Postgres; without it, attendance would vanish on the next restart."
        )
    init_db()


def utcnow():
    return datetime.now(timezone.utc)


def local(dt) -> str:
    """UTC on disk, Bogota on screen."""
    if not dt:
        return ""
    return dt.astimezone(LOCAL_TZ).strftime("%Y-%m-%d %H:%M")


def local_time(dt) -> str:
    if not dt:
        return ""
    return dt.astimezone(LOCAL_TZ).strftime("%H:%M:%S")


def e(text) -> str:
    """
    Escape everything that reaches a page.

    Version 1 put roster names straight into HTML. That was tolerable while
    the professor wrote the roster; now students type their own names, and a
    "name" containing a <script> tag would run on the professor's screen.
    """
    return html.escape("" if text is None else str(text))


# ---------------------------------------------------------------
# BROWSER, SYSTEM AND APP
#
# Three words per check-in, read from the User-Agent header. The full header
# is never stored: combined with other details it can help single out one
# phone, which goes against the privacy design of the project.
#
# Detection is best effort. WhatsApp on iPhone opens links in a Safari view
# that looks exactly like Safari, so it cannot be told apart.
# ---------------------------------------------------------------

def parse_ua(ua: str) -> dict:
    ua = ua or ""
    if re.search(r"iPhone|iPad|iPod", ua):
        system = "iOS"
    elif "Android" in ua:
        system = "Android"
    elif "Windows" in ua:
        system = "Windows"
    elif "Macintosh" in ua or "Mac OS X" in ua:
        system = "macOS"
    elif "Linux" in ua:
        system = "Linux"
    else:
        system = "Otro"

    in_app = ""
    if "Instagram" in ua:
        in_app = "Instagram"
    elif re.search(r"FBAN|FBAV|FB_IAB", ua):
        in_app = "Facebook"
    elif "WhatsApp" in ua:
        in_app = "WhatsApp"
    elif "; wv)" in ua:
        in_app = "App (WebView)"

    if "SamsungBrowser" in ua:
        browser = "Samsung Internet"
    elif re.search(r"EdgA|EdgiOS|Edg/", ua):
        browser = "Edge"
    elif re.search(r"FxiOS|Firefox", ua):
        browser = "Firefox"
    elif re.search(r"CriOS|Chrome", ua):
        browser = "Chrome"
    elif "Safari" in ua:
        browser = "Safari"
    else:
        browser = "Otro"

    return {"browser": browser, "os": system, "in_app": in_app}


def ua_of(request: Request) -> dict:
    return parse_ua(request.headers.get("user-agent", ""))


def ua_text(ua: dict) -> str:
    text = f"{ua['browser']} · {ua['os']}"
    return text + (f" · desde {ua['in_app']}" if ua.get("in_app") else "")


# ---------------------------------------------------------------
# EVENTS
# ---------------------------------------------------------------

EVENT_LABELS = {
    "vinculacion": "Vinculó este celular durante la clase",
    "cambio_dispositivo": "Su carné pasó a otro navegador o celular durante la clase",
    "rafaga": "Varios celulares nuevos seguidos con el mismo navegador y sistema",
    "desde_app": "Abrió el enlace dentro de una app",
    "carne_no_en_lista": "Marcó con un carné o usuario que no está en la lista",
    "qr_vencido": "Escaneó un QR vencido",
    "manual": "Marcado a mano por el profesor",
    "desvinculado": "El profesor desvinculó su celular",
    "carne_corregido": "El profesor corrigió su carné o su nombre",
    "validado_con_lista": "El profesor lo validó con la lista",
}


def log_event(conn, session_id, student_id, kind, detail="", review=False, ua=None):
    ua = ua or {}
    conn.execute(
        "INSERT INTO events (session_id, student_id, kind, detail, review, "
        "browser, os, in_app) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
        (session_id, student_id, kind, detail, review,
         ua.get("browser"), ua.get("os"), ua.get("in_app")),
    )


def check_burst(conn, session_id, ua):
    """
    Three or more new links within a minute, all from the same browser and
    system, during one class.

    On its own this proves nothing: half a class can own the same iPhone
    model. It is the shape someone typing classmates' carnes into fresh
    private windows would leave behind, so it goes under "Por revisar".
    """
    since = utcnow() - timedelta(seconds=BURST_WINDOW_SECONDS)
    n = conn.execute(
        "SELECT COUNT(*) AS n FROM events WHERE session_id = %s "
        "AND kind = 'vinculacion' AND created_at >= %s "
        "AND browser = %s AND os = %s",
        (session_id, since, ua["browser"], ua["os"]),
    ).fetchone()["n"]
    if n < BURST_COUNT:
        return
    already = conn.execute(
        "SELECT 1 FROM events WHERE session_id = %s AND kind = 'rafaga' "
        "AND created_at >= %s AND browser = %s AND os = %s",
        (session_id, since, ua["browser"], ua["os"]),
    ).fetchone()
    if not already:
        names = [r["name"] for r in conn.execute(
            "SELECT st.name FROM events ev JOIN students st ON st.id = ev.student_id "
            "WHERE ev.session_id = %s AND ev.kind = 'vinculacion' AND ev.created_at >= %s "
            "AND ev.browser = %s AND ev.os = %s ORDER BY ev.created_at",
            (session_id, since, ua["browser"], ua["os"])).fetchall()]
        log_event(conn, session_id, None, "rafaga",
                  f"{n} celulares nuevos en menos de {BURST_WINDOW_SECONDS} s con "
                  f"{ua['browser']} en {ua['os']}: {', '.join(names[:6])}. En la primera "
                  f"clase es normal; revisa solo si no reconoces a alguien o no estaba", True, ua)


# ---------------------------------------------------------------
# LINKING A CARNE TO A BROWSER
# ---------------------------------------------------------------

class RetryNeeded(Exception):
    """Someone changed the same row at the same instant."""


def bind_device(conn, student, device_id: str, session_id):
    """
    Point a carne at this browser.

    Returns:
        "ok"     first link for this carne
        "moved"  taken from another browser, written to device_changes
        "same"   already this browser
        "taken"  this browser already belongs to another student
    Raises RetryNeeded if the row changed while we were working.
    """
    owner = conn.execute(
        "SELECT id FROM students WHERE device_id = %s", (device_id,)
    ).fetchone()
    if owner and owner["id"] != student["id"]:
        # The one rule that does not bend: a browser marks for one person.
        return "taken"

    previous = student["device_id"]
    if previous == device_id:
        return "same"

    # The WHERE clause carries the value we read, NULL included, so the write
    # only lands if nothing changed since. Two people linking the same carne
    # at the same instant still produce exactly one winner.
    cur = conn.execute(
        "UPDATE students SET device_id = %s, device_no = device_no + 1, "
        "registered_at = COALESCE(registered_at, now()) "
        "WHERE id = %s AND device_id IS NOT DISTINCT FROM %s",
        (device_id, student["id"], previous),
    )
    if cur.rowcount == 0:
        raise RetryNeeded()

    if previous:
        conn.execute(
            "INSERT INTO device_changes (student_id, old_device, new_device, session_id) "
            "VALUES (%s, %s, %s, %s)",
            (student["id"], previous, device_id, session_id),
        )
        return "moved"
    return "ok"


def make_receipt(conn, course_id: int) -> str:
    day = utcnow().astimezone(LOCAL_TZ).strftime("%m%d")
    while True:
        tail = "".join(secrets.choice(RECEIPT_ALPHABET) for _ in range(4))
        code = f"C{course_id}-{day}-{tail}"
        if not conn.execute("SELECT 1 FROM attendance WHERE receipt = %s",
                            (code,)).fetchone():
            return code


def record_attendance(conn, session, student, method, ua=None):
    """
    Returns (attendance_row, created). A second check-in by the same student
    in the same class creates nothing and returns the first one.
    """
    ua = ua or {}
    row = conn.execute(
        "INSERT INTO attendance (session_id, student_id, method, receipt, "
        "browser, os, in_app, device_no) "
        "VALUES (%s, %s, %s, %s, %s, %s, %s, %s) "
        "ON CONFLICT (session_id, student_id) DO NOTHING RETURNING *",
        (session["id"], student["id"], method,
         make_receipt(conn, session["course_id"]),
         ua.get("browser"), ua.get("os"), ua.get("in_app"),
         student.get("device_no") if method == "qr" else None),
    ).fetchone()
    if row:
        return row, True
    row = conn.execute(
        "SELECT * FROM attendance WHERE session_id = %s AND student_id = %s",
        (session["id"], student["id"]),
    ).fetchone()
    return row, False


def enrollment(conn, course_id, student_id):
    return conn.execute(
        "SELECT * FROM enrollments WHERE course_id = %s AND student_id = %s",
        (course_id, student_id),
    ).fetchone()


def normalize_carne(value) -> str:
    """Excel turns 2021104567 into 2021104567.0 and likes leading apostrophes."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    # Excel may add a leading apostrophe, and Brightspace exports write
    # "#agomez" and "#170001" so that Excel keeps them as text.
    text = str(value or "").strip().lstrip("'#")
    # A student may type the whole email; only the part before @ is the key.
    text = text.split("@")[0]
    return re.sub(r"\s+", "", text).upper()


def find_student(conn, key: str):
    """The same person can be found by carne or by institutional username."""
    return conn.execute(
        "SELECT * FROM students WHERE carne = %s OR username = %s "
        "ORDER BY (username = %s) DESC NULLS LAST LIMIT 1",
        (key, key, key),
    ).fetchone()


# ---------------------------------------------------------------
# TEACHER ACCOUNTS
#
# Version 1 had one shared password, so nobody could tell which person had
# marked a student by hand. Each professor now has an account.
# ---------------------------------------------------------------

class NeedLogin(Exception):
    pass


@app.exception_handler(NeedLogin)
def need_login(request: Request, exc: NeedLogin):
    if request.url.path.endswith("/state"):
        return JSONResponse({"error": "login"}, status_code=401)
    return RedirectResponse("/login", status_code=303)


@app.exception_handler(RetryNeeded)
def retry_needed(request: Request, exc: RetryNeeded):
    return HTMLResponse(student_page(
        "<p class='bad-title'>Inténtalo otra vez</p>"
        "<p>Alguien usó ese carné en el mismo instante. Vuelve a escanear el QR.</p>"
    ))


@app.exception_handler(psycopg.errors.UniqueViolation)
def unique_violation(request: Request, exc):
    return HTMLResponse(student_page(
        "<p class='bad-title'>Inténtalo otra vez</p>"
        "<p>Dos registros chocaron al mismo tiempo. Vuelve a escanear el QR.</p>"
    ))


def current_teacher(request: Request, conn):
    token = request.cookies.get("teacher_session")
    if not token:
        return None
    return conn.execute(
        "SELECT t.* FROM teacher_sessions s JOIN teachers t ON t.id = s.teacher_id "
        "WHERE s.token = %s AND s.expires_at > now()",
        (token,),
    ).fetchone()


def require_teacher(request: Request, conn):
    teacher = current_teacher(request, conn)
    if not teacher:
        raise NeedLogin()
    return teacher


def own_course(conn, teacher, course_id):
    course = conn.execute(
        "SELECT * FROM courses WHERE id = %s AND teacher_id = %s",
        (course_id, teacher["id"]),
    ).fetchone()
    if not course:
        # Another professor's class looks exactly like one that does not exist.
        raise NeedLogin()
    return course


def own_session(conn, teacher, session_id):
    session = conn.execute(
        "SELECT s.*, c.name AS course_name, c.has_roster, c.teacher_id "
        "FROM sessions s JOIN courses c ON c.id = s.course_id "
        "WHERE s.id = %s AND c.teacher_id = %s",
        (session_id, teacher["id"]),
    ).fetchone()
    if not session:
        raise NeedLogin()
    return session


def start_teacher_session(conn, teacher_id) -> Response:
    token = secrets.token_urlsafe(32)
    conn.execute(
        "INSERT INTO teacher_sessions (token, teacher_id, expires_at) VALUES (%s, %s, %s)",
        (token, teacher_id, utcnow() + timedelta(hours=TEACHER_SESSION_HOURS)),
    )
    response = RedirectResponse("/", status_code=303)
    response.set_cookie("teacher_session", token,
                        max_age=TEACHER_SESSION_HOURS * 3600,
                        httponly=True, samesite="lax")
    return response


# A real bcrypt hash of a random string. Checking a wrong email against it
# takes as long as checking a real one, so response time does not reveal
# which emails have accounts.
_DUMMY_HASH = bcrypt.hashpw(b"unipresence-dummy", bcrypt.gensalt()).decode()


# ---------------------------------------------------------------
# PAGE SHELLS
# ---------------------------------------------------------------

BASE_CSS = """
:root { --bg:#f2f5f7; --card:#fff; --ink:#15212e; --muted:#5a6876; --line:#d6dde4;
        --ok:#18794e; --ok-soft:#e3f4ec; --bad:#b42318; --bad-soft:#fdecea;
        --warn:#a15c07; --warn-soft:#fdf3e1; --accent:#1f5f8b; }
* { box-sizing:border-box; }
body { margin:0; background:var(--bg); color:var(--ink);
       font:16px/1.5 -apple-system, "Segoe UI", Roboto, sans-serif; }
a { color:var(--accent); }
.wrap { max-width:1000px; margin:0 auto; padding:20px 16px 60px; }
header.bar { background:var(--card); border-bottom:1px solid var(--line); }
header.bar .wrap { padding-block:12px; display:flex; gap:16px; align-items:center;
                   justify-content:space-between; flex-wrap:wrap; }
.brand { font-weight:700; font-size:18px; text-decoration:none; color:var(--ink); }
nav { display:flex; gap:16px; align-items:center; font-size:15px; color:var(--muted); }
h1 { font-size:26px; margin:8px 0 16px; }
h2 { font-size:19px; margin:28px 0 10px; }
.card { background:var(--card); border:1px solid var(--line); border-radius:10px; padding:18px; }
.grid { display:grid; grid-template-columns:repeat(auto-fill, minmax(260px, 1fr)); gap:14px; }
.btn { display:inline-block; background:var(--ok); color:#fff; border:0; border-radius:8px;
       padding:11px 18px; font-size:16px; font-weight:600; text-decoration:none; cursor:pointer; }
.btn.big { padding:16px 24px; font-size:19px; width:100%; text-align:center; }
.btn.grey { background:#5a6876; }
.btn.red { background:var(--bad); }
.btn.light { background:var(--ok-soft); color:var(--ok); }
.btn.small { padding:6px 12px; font-size:14px; }
input, textarea, select { font:inherit; width:100%; padding:11px 12px; border:1px solid var(--line);
       border-radius:8px; background:#fff; color:var(--ink); }
textarea { min-height:150px; font-family:ui-monospace, Menlo, monospace; font-size:14px; }
label { display:block; font-weight:600; margin:14px 0 6px; }
.hint { color:var(--muted); font-size:14px; margin:4px 0 0; }
.msg { padding:12px 14px; border-radius:8px; margin:12px 0; }
.msg.bad { background:var(--bad-soft); color:var(--bad); }
.msg.ok { background:var(--ok-soft); color:var(--ok); }
.msg.warn { background:var(--warn-soft); color:var(--warn); }
.table-wrap { overflow-x:auto; }
table { width:100%; border-collapse:collapse; background:var(--card); font-size:14px; }
th, td { text-align:left; padding:8px 10px; border-bottom:1px solid var(--line); vertical-align:top; }
th { color:var(--muted); font-weight:600; font-size:12px; text-transform:uppercase; letter-spacing:.04em; }
.pill { display:inline-block; padding:2px 8px; border-radius:999px; font-size:12px; font-weight:600; }
.pill.ok { background:var(--ok-soft); color:var(--ok); }
.pill.bad { background:var(--bad-soft); color:var(--bad); }
.pill.warn { background:var(--warn-soft); color:var(--warn); }
.pill.grey { background:#e8edf1; color:var(--muted); }
.muted { color:var(--muted); }
form.inline { display:inline; }
"""


def teacher_page(title: str, body: str, teacher=None) -> str:
    nav = ""
    if teacher:
        nav = (f"<nav><a href='/'>Mis clases</a><span>{e(teacher['name'])}</span>"
               f"<a href='/logout'>Salir</a></nav>")
    return f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(title)} · UniPresence</title><style>{BASE_CSS}</style></head>
<body><header class="bar"><div class="wrap"><a class="brand" href="/">UniPresence</a>{nav}</div></header>
<main class="wrap">{body}</main></body></html>"""


STUDENT_CSS = """
body { background:var(--bg); }
.s { max-width:440px; margin:0 auto; padding:28px 18px 48px; text-align:center; }
.course { font-size:26px; font-weight:700; margin:6px 0 2px; }
.prof { color:var(--muted); margin:0 0 18px; }
.name { font-size:28px; font-weight:700; margin:14px 0 2px; }
.carne { color:var(--muted); margin:0 0 20px; }
.s form { text-align:left; }
.bad-title { color:var(--bad); font-size:24px; font-weight:700; margin:24px 0 8px; }
.okscreen { position:fixed; inset:0; background:var(--ok); color:#fff; display:flex;
            align-items:center; justify-content:center; text-align:center; padding:24px; }
.okscreen .inner { max-width:440px; width:100%; }
.check { width:96px; height:96px; border-radius:50%; background:#fff; color:var(--ok);
         font-size:60px; line-height:96px; margin:0 auto 14px; font-weight:700; }
.oktitle { font-size:30px; font-weight:800; margin:0 0 14px; }
.okscreen .name { color:#fff; }
.okline { font-size:17px; margin:2px 0; opacity:.95; }
.clock { font-size:44px; font-weight:800; font-variant-numeric:tabular-nums; margin:16px 0 4px; }
.live { display:inline-flex; align-items:center; gap:8px; font-size:14px; opacity:.9; }
.dot { width:10px; height:10px; border-radius:50%; background:#fff; animation:pulse 1s infinite; }
@keyframes pulse { 50% { opacity:.25; } }
.receipt { margin-top:18px; background:rgba(255,255,255,.16); border-radius:10px; padding:12px; }
.receipt b { font-size:24px; letter-spacing:2px; font-family:ui-monospace, Menlo, monospace; }
"""


def student_page(body: str) -> str:
    return f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>UniPresence</title><style>{BASE_CSS}{STUDENT_CSS}</style></head>
<body><div class="s">{body}</div></body></html>"""


# ---------------------------------------------------------------
# SIGN UP, LOG IN
# ---------------------------------------------------------------

PRIVACY_CONTACT = os.environ.get("PRIVACY_CONTACT", "al profesor que toma la asistencia")


@app.get("/privacidad", response_class=HTMLResponse)
def privacy_page():
    """The data policy a person accepts the first time they link a phone."""
    return student_page(f"""<div style="text-align:left">
      <h1 style="font-size:22px">Política de tratamiento de datos</h1>
      <p>UniPresence es un prototipo para registrar asistencia. Se usa con
      autorización previa, según la Ley 1581 de 2012.</p>
      <p><strong>Qué se guarda:</strong> tu carné, tu usuario institucional (lo que va
      antes de @, nunca el correo completo) o la identificación que escribas; tu nombre;
      la hora en que marcas; el navegador, el sistema del celular y si abriste el enlace
      dentro de una app. Además, un código al azar en una cookie para reconocer este
      navegador.</p>
      <p><strong>Qué no se guarda:</strong> ubicación, fotos, datos biométricos ni el número
      del celular.</p>
      <p><strong>Para qué:</strong> solo para la asistencia de la clase o evento, y para
      detectar si alguien marca por otra persona.</p>
      <p><strong>Quién lo ve:</strong> solo el profesor u organizador de esa clase o evento.</p>
      <p><strong>Tus derechos:</strong> conocer, actualizar, corregir o pedir que se borren tus
      datos. Pídeselo {e(PRIVACY_CONTACT)}.</p>
      <p class="hint">Si no aceptas, no se guarda nada y el profesor puede marcarte a mano.</p>
    </div>""")


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/login", response_class=HTMLResponse)
def login_page(error: str = ""):
    msg = "<div class='msg bad'>Correo o contraseña incorrectos.</div>" if error else ""
    return teacher_page("Entrar", f"""
      <div class="card" style="max-width:420px;margin:30px auto">
        <h1>Entrar</h1>{msg}
        <form method="post" action="/login">
          <label for="email">Correo</label>
          <input id="email" name="email" type="email" required autofocus>
          <label for="password">Contraseña</label>
          <input id="password" name="password" type="password" required>
          <p style="margin-top:18px"><button class="btn big" type="submit">Entrar</button></p>
        </form>
        <p class="hint">¿Primera vez? <a href="/signup">Crea tu cuenta de profesor</a>.</p>
      </div>""")


@app.post("/login")
def do_login(email: str = Form(...), password: str = Form(...)):
    with db() as conn:
        teacher = conn.execute("SELECT * FROM teachers WHERE email = %s",
                               (email.strip().lower(),)).fetchone()
        stored = teacher["password_hash"] if teacher else _DUMMY_HASH
        valid = bcrypt.checkpw(password.encode(), stored.encode())
        if not teacher or not valid:
            return RedirectResponse("/login?error=1", status_code=303)
        return start_teacher_session(conn, teacher["id"])


@app.get("/signup", response_class=HTMLResponse)
def signup_page(error: str = ""):
    msg = f"<div class='msg bad'>{e(error)}</div>" if error else ""
    code_field = ""
    if SIGNUP_CODE:
        code_field = """<label for="code">Código de invitación</label>
          <input id="code" name="code" required>"""
    return teacher_page("Crear cuenta", f"""
      <div class="card" style="max-width:460px;margin:30px auto">
        <h1>Cuenta de profesor</h1>{msg}
        <form method="post" action="/signup">
          <label for="name">Nombre</label>
          <input id="name" name="name" required maxlength="80" placeholder="Prof. María Pérez">
          <label for="institution">Institución <span class="muted">(opcional)</span></label>
          <input id="institution" name="institution" maxlength="120" placeholder="Universidad">
          <label for="email">Correo</label>
          <input id="email" name="email" type="email" required>
          <label for="password">Contraseña</label>
          <input id="password" name="password" type="password" required minlength="8">
          <p class="hint">Mínimo 8 caracteres. No uses la misma de la universidad.</p>
          {code_field}
          <p style="margin-top:18px"><button class="btn big" type="submit">Crear cuenta</button></p>
        </form>
        <p class="hint">¿Ya tienes cuenta? <a href="/login">Entra aquí</a>.</p>
      </div>""")


@app.post("/signup")
def do_signup(name: str = Form(...), email: str = Form(...), password: str = Form(...),
              institution: str = Form(""), code: str = Form("")):
    email = email.strip().lower()
    if SIGNUP_CODE and not secrets.compare_digest(code.strip(), SIGNUP_CODE):
        return RedirectResponse("/signup?error=Código+de+invitación+incorrecto", status_code=303)
    if len(password) < 8:
        return RedirectResponse("/signup?error=La+contraseña+necesita+8+caracteres", status_code=303)
    hashed = bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()
    with db() as conn:
        row = conn.execute(
            "INSERT INTO teachers (email, name, institution, password_hash) "
            "VALUES (%s, %s, %s, %s) ON CONFLICT (email) DO NOTHING RETURNING id",
            (email, name.strip()[:80], institution.strip()[:120], hashed),
        ).fetchone()
        if not row:
            return RedirectResponse("/signup?error=Ya+existe+una+cuenta+con+ese+correo", status_code=303)
        return start_teacher_session(conn, row["id"])


@app.get("/logout")
def logout(request: Request):
    token = request.cookies.get("teacher_session")
    if token:
        with db() as conn:
            conn.execute("DELETE FROM teacher_sessions WHERE token = %s", (token,))
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie("teacher_session")
    return response


# ---------------------------------------------------------------
# CLASSES
# ---------------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
def home(request: Request, msg: str = ""):
    with db() as conn:
        teacher = require_teacher(request, conn)
        courses = conn.execute("""
            SELECT c.*,
                   (SELECT COUNT(*) FROM enrollments en WHERE en.course_id = c.id) AS students,
                   (SELECT COUNT(*) FROM enrollments en WHERE en.course_id = c.id
                       AND en.status = 'pendiente') AS pending,
                   (SELECT id FROM sessions s WHERE s.course_id = c.id AND s.active
                       ORDER BY id DESC LIMIT 1) AS live_session
            FROM courses c WHERE c.teacher_id = %s
            ORDER BY c.last_used_at DESC NULLS LAST, c.created_at DESC
        """, (teacher["id"],)).fetchall()

    cards = []
    for c in courses:
        list_note = "con lista" if c["has_roster"] else "sin lista: se arma con quienes marquen"
        pending = (f" · <span class='pill warn'>{c['pending']} por aprobar</span>"
                   if c["pending"] else "")
        if c["live_session"]:
            action = (f"<a class='btn big' href='/session/{c['live_session']}/live'>"
                      f"Volver a la asistencia en curso</a>")
        else:
            action = (f"<form method='post' action='/course/{c['id']}/start'>"
                      f"<button class='btn big' type='submit'>Tomar asistencia</button></form>")
        cards.append(f"""
          <div class="card">
            <h2 style="margin-top:0">{e(c['name'])}</h2>
            <p class="muted">{c['students']} estudiantes · {list_note}{pending}</p>
            {action}
            <p style="margin:12px 0 0;display:flex;gap:18px"><a href="/course/{c['id']}">Ver lista</a>
              <a href="/course/{c['id']}/historial">Historial</a></p>
          </div>""")

    empty = ("<div class='card'><p>Todavía no tienes clases. Crea la primera: "
             "con lista o sin lista, las dos funcionan.</p></div>")
    flash = f"<div class='msg ok'>{e(msg)}</div>" if msg else ""
    return teacher_page("Mis clases", f"""
      <div style="display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap">
        <h1>Mis clases</h1><a class="btn" href="/course/new">Nueva clase</a>
      </div>{flash}
      <div class="grid">{''.join(cards) or empty}</div>""", teacher)


ROSTER_HELP = """
<p class="hint">Opcional. Pega desde Excel dos columnas, nombre y carné, o sube el
archivo (.xlsx o .csv). La exportación de calificaciones de Brightspace sirve tal
como sale: si trae varios grupos, te preguntamos cuáles importar. Si no tienes la lista, déjalo vacío: la lista se arma sola
con los estudiantes que marquen, y tú los apruebas.</p>
"""


@app.get("/course/new", response_class=HTMLResponse)
def new_course_page(request: Request):
    with db() as conn:
        teacher = require_teacher(request, conn)
    return teacher_page("Nueva clase", f"""
      <h1>Nueva clase</h1>
      <div class="card" style="max-width:640px">
        <form method="post" action="/course/new" enctype="multipart/form-data">
          <label for="name">Nombre de la clase</label>
          <input id="name" name="name" required maxlength="100" placeholder="Genética · Grupo 1">
          <label for="roster_text">Lista de estudiantes</label>
          {ROSTER_HELP}
          <textarea id="roster_text" name="roster_text" placeholder="María Fernanda Rojas&#9;2021104567&#10;Juan Sebastián Pérez&#9;2021104568"></textarea>
          <label for="roster_file">O sube el archivo</label>
          <input id="roster_file" name="roster_file" type="file" accept=".xlsx,.csv,.txt">
          <p style="margin-top:18px"><button class="btn big" type="submit">Crear clase</button></p>
        </form>
      </div>""", teacher)


def _sections(rows):
    counts = {}
    for r in rows:
        section = r[2]
        if section:
            counts[section] = counts.get(section, 0) + 1
    return dict(sorted(counts.items()))


def short_section(label: str) -> str:
    """Brightspace writes 'SISTEMA NERVIOSO2026303614'; people say 'grupo 3614'."""
    m = re.search(r"(\d{4})\s*$", label or "")
    return f"Grupo {m.group(1)}" if m else (label or "Sin grupo")


def section_picker(teacher, action: str, name: str, rows, problems) -> HTMLResponse:
    """
    A platform export often holds every group of a subject: the Brightspace
    file for Fundamentos had 188 students in 9 groups. The professor keeps the
    groups that are hers, all of them by default; which ones sit in the room
    on a given day is chosen later, each time attendance starts.

    The parsed list travels in a hidden field to the next step. It is the
    professor's own data going back to the professor, so nothing is gained by
    storing it in between.
    """
    boxes = "".join(
        f"<label style='font-weight:normal;display:flex;gap:10px;align-items:center'>"
        f"<input type='checkbox' name='sections' value='{e(sec)}' checked style='width:auto'>"
        f"<span><strong>{e(short_section(sec))}</strong> <span class='muted'>· {n} estudiantes · {e(sec)}</span></span></label>"
        for sec, n in _sections(rows).items())
    warn = (f"<div class='msg warn'>Revisa: {e(' '.join(problems[:5]))}</div>" if problems else "")
    return HTMLResponse(teacher_page("Elige los grupos", f"""
      <h1>¿Qué grupos son tuyos?</h1>
      <p>El archivo trae {len(rows)} estudiantes en {len(_sections(rows))} grupos.
      Deja marcados los grupos que dictas. Antes de cada clase vas a escoger cuáles
      están en el salón ese día.</p>{warn}
      <div class="card" style="max-width:640px">
        <form method="post" action="{action}">
          <input type="hidden" name="name" value="{e(name)}">
          <input type="hidden" name="payload" value="{e(json.dumps(rows))}">
          <div style="display:grid;gap:8px">{boxes}</div>
          <p style="margin-top:18px"><button class="btn big" type="submit">Guardar lista</button></p>
        </form>
      </div>""", teacher))


def _picked(payload: str, sections: list[str]):
    try:
        rows = [tuple(r) for r in json.loads(payload)]
    except (ValueError, TypeError):
        return []
    chosen = set(sections)
    return [r for r in rows if len(r) in (3, 4, 5) and r[2] in chosen]


def _create_course(conn, teacher, name, rows):
    course = conn.execute(
        "INSERT INTO courses (teacher_id, name) VALUES (%s, %s) RETURNING *",
        (teacher["id"], name.strip()[:100] or "Clase"),
    ).fetchone()
    added = save_roster(conn, course["id"], rows) if rows else 0
    return course, added


@app.post("/course/new")
async def create_course(request: Request, name: str = Form(...), roster_text: str = Form(""),
                        roster_file: UploadFile | None = File(None)):
    rows, problems = await read_roster_input(roster_text, roster_file)
    with db() as conn:
        teacher = require_teacher(request, conn)
        if len(_sections(rows)) > 1:
            return section_picker(teacher, "/course/new/groups", name, rows, problems)
        course, added = _create_course(conn, teacher, name, rows)
    note = f"Clase creada con {added} estudiantes." if added else "Clase creada sin lista."
    if problems:
        note += " Revisa: " + " ".join(problems[:5])
    return RedirectResponse(f"/course/{course['id']}?msg={note}", status_code=303)


@app.post("/course/new/groups")
def create_course_groups(request: Request, name: str = Form(...), payload: str = Form(...),
                         sections: list[str] = Form(default=[])):
    rows = _picked(payload, sections)
    with db() as conn:
        teacher = require_teacher(request, conn)
        course, added = _create_course(conn, teacher, name, rows)
    note = (f"Clase creada con {added} estudiantes de {len(sections)} grupo(s)." if added
            else "Clase creada sin lista: no marcaste ningún grupo.")
    return RedirectResponse(f"/course/{course['id']}?msg={note}", status_code=303)


@app.get("/course/{course_id}/historial", response_class=HTMLResponse)
def course_history(course_id: int, request: Request):
    """The attendance already taken in a class, on its own page."""
    return course_page(course_id, request, only_history=True)


@app.get("/course/{course_id}", response_class=HTMLResponse)
def course_page(course_id: int, request: Request, msg: str = "", only_history: bool = False):
    with db() as conn:
        teacher = require_teacher(request, conn)
        course = own_course(conn, teacher, course_id)
        students = conn.execute("""
            SELECT st.id, st.carne, st.name, st.device_id IS NOT NULL AS linked, en.status,
                   en.section
            FROM enrollments en JOIN students st ON st.id = en.student_id
            WHERE en.course_id = %s ORDER BY en.status DESC, en.section, st.name
        """, (course_id,)).fetchall()
        sessions = conn.execute("""
            SELECT s.*, (SELECT COUNT(*) FROM attendance a WHERE a.session_id = s.id) AS present
            FROM sessions s WHERE s.course_id = %s ORDER BY s.id DESC
        """, (course_id,)).fetchall()

    pending = [s for s in students if s["status"] == "pendiente"]
    group_counts = {}
    for s_ in students:
        if s_["status"] == "aprobado" and s_["section"]:
            group_counts[s_["section"]] = group_counts.get(s_["section"], 0) + 1
    groups_html, start_groups = "", ""
    if len(group_counts) > 1:
        # Choosing today's groups happens here, next to the button, in one step.
        # Nothing checked means every group, so a quick start never fails.
        start_groups = (
            "<input type='hidden' name='chosen' value='1'>"
            "<p style='margin:0 0 6px'><strong>¿Qué grupos están hoy en el salón?</strong> "
            "<span class='hint'>Si no marcas ninguno, se cuentan todos.</span></p>"
            "<div style='display:flex;flex-wrap:wrap;gap:8px;margin-bottom:12px'>" + "".join(
                f"<label style='font-weight:normal;display:flex;gap:6px;align-items:center;margin:0;"
                f"padding:6px 10px;border:1px solid var(--line);border-radius:8px;background:#fff'>"
                f"<input type='checkbox' name='sections' value='{e(g)}' style='width:auto'>"
                f"{e(short_section(g))} <span class='muted'>· {n}</span></label>"
                for g, n in sorted(group_counts.items())) + "</div>")
        items = "".join(
            f"<form class='inline' method='post' action='/course/{course_id}/groups/remove' "
            f"onsubmit=\"return confirm('¿Quitar {e(short_section(g))} de esta clase? "
            f"Los que ya tienen asistencia se quedan.')\" style='margin:0 8px 6px 0'>"
            f"<input type='hidden' name='section' value='{e(g)}'>"
            f"<span class='pill ok'>{e(short_section(g))} · {n}</span> "
            f"<button class='btn small grey' type='submit'>Quitar</button></form>"
            for g, n in sorted(group_counts.items()))
        groups_html = (f"<details style='margin-top:12px'><summary>Quitar un grupo de esta clase</summary>"
                       f"<div class='card' style='margin-top:10px'><p class='hint'>Solo si subiste un grupo "
                       f"que no es de esta clase. Quitar no es para tomar asistencia: para eso marca los "
                       f"grupos de hoy arriba, junto a “Tomar asistencia”. Los estudiantes que ya tienen "
                       f"asistencia no se quitan.</p><div style='display:flex;flex-wrap:wrap'>{items}</div></div></details>")
    pending_html = ""
    with db() as conn:
        repeated = sum(1 for p, m, _ in suggest_matches(conn, course_id)
                       if m and p["status"] == "aprobado")
    if repeated and not pending:
        pending_html = f"""
          <div class="msg warn" style="margin-top:12px">{repeated} estudiantes aprobados parecen estar
          dos veces: una como se escribieron al marcar y otra en la lista.
          <a class="btn small" href="/course/{course_id}/validar">Validar con la lista</a></div>"""
    validate_link = (f"<a class='btn small' href='/course/{course_id}/validar'>Validar con la lista</a>"
                     if pending and any(s["status"] == "aprobado" for s in students) else
                     "<span class='hint'>¿Tomaste asistencia sin lista? Sube la lista abajo, en "
                     "“Agregar o actualizar la lista”, y la app los cruza.</span>")
    if pending:
        rows = "".join(
            f"<tr><td>{e(s['name'])}</td><td>{e(s['carne'])}</td><td>"
            f"<form class='inline' method='post' action='/course/{course_id}/approve'>"
            f"<input type='hidden' name='student_id' value='{s['id']}'>"
            f"<button class='btn small' type='submit'>Aprobar</button></form> "
            f"<form class='inline' method='post' action='/course/{course_id}/remove'>"
            f"<input type='hidden' name='student_id' value='{s['id']}'>"
            f"<button class='btn small grey' type='submit'>Quitar</button></form> "
            f"<a href='/course/{course_id}/student/{s['id']}/edit'>Corregir</a></td></tr>"
            for s in pending)
        pending_html = f"""
          <h2>Por aprobar ({len(pending)})</h2>
          <p class="hint">Se unieron marcando asistencia. Si la clase tiene lista, usa primero
          “Validar con la lista”: une a cada uno con su fila. “Aprobar” es solo para quien no está en la lista.</p>
          <form method="post" action="/course/{course_id}/approve" class="inline">
            <input type="hidden" name="all" value="1">
            <button class="btn light small" type="submit">Aprobar a todos</button></form>
          {validate_link}
          <div class="table-wrap" style="margin-top:10px"><table>
            <tr><th>Nombre</th><th>Carné</th><th></th></tr>{rows}</table></div>"""

    approved = [s for s in students if s["status"] == "aprobado"]
    linked_yes = "<span class='pill ok'>sí</span>"
    linked_no = "<span class='muted'>todavía no</span>"
    def unlink_button(st):
        if not st["linked"]:
            return ""
        return (f" <form class='inline' method='post' action='/course/{course_id}/unlink' "
                f"onsubmit=\"return confirm('¿Desvincular el celular de {e(st['name'])}? "
                f"La próxima vez tendrá que escribir su carné.')\">"
                f"<input type='hidden' name='student_id' value='{st['id']}'>"
                f"<button class='btn small grey' type='submit'>Desvincular</button></form>")

    roster_rows = "".join(
        f"<tr><td>{e(s['name'])}</td><td>{e(s['carne'])}</td>"
        f"<td>{e(short_section(s['section'])) if s['section'] else ''}</td>"
        f"<td>{linked_yes if s['linked'] else linked_no}{unlink_button(s)}</td>"
        f"<td><a href='/course/{course_id}/student/{s['id']}/edit'>Corregir</a></td></tr>"
        for s in approved)

    live_pill = "<span class='pill ok'>en curso</span>"
    session_rows = ""
    for s in sessions:
        ended = live_pill if s["active"] else local(s["ended_at"])
        links = f"<a href='/session/{s['id']}/report'>Reporte</a>"
        if s["active"]:
            links = f"<a href='/session/{s['id']}/live'>Pantalla en vivo</a> · " + links
        else:
            links += (f" · <a href='/session/{s['id']}/delete' style='color:var(--bad)'>Borrar</a>")
        session_rows += (f"<tr><td>#{s['id']}</td><td>{local(s['started_at'])}</td>"
                         f"<td>{ended}</td><td>{s['present']}</td><td>{links}</td></tr>")

    flash = f"<div class='msg ok'>{e(msg)}</div>" if msg else ""
    if only_history:
        return teacher_page(f"Historial · {course['name']}", f"""
      <h1>{e(course['name'])} · Historial</h1>
      <p><a href="/course/{course_id}">Ver lista y tomar asistencia</a></p>
      <div class="table-wrap"><table>
        <tr><th>Sesión</th><th>Inicio</th><th>Fin</th><th>Presentes</th><th></th></tr>
        {session_rows or "<tr><td colspan=5 class=muted>Todavía no hay sesiones.</td></tr>"}
      </table></div>""", teacher)
    return teacher_page(course["name"], f"""
      <h1>{e(course['name'])}</h1>{flash}
      <form method="post" action="/course/{course_id}/start" style="max-width:640px">
        {start_groups}
        <button class="btn big" type="submit" style="max-width:420px">Tomar asistencia</button></form>
      {pending_html}
      <h2>Lista ({len(approved)})</h2>
      <div class="table-wrap"><table>
        <tr><th>Nombre</th><th>Carné</th><th>Grupo</th><th>Celular vinculado</th><th></th></tr>
        {roster_rows or "<tr><td colspan=5 class=muted>Sin estudiantes todavía.</td></tr>"}
      </table></div>
      <details style="margin-top:12px"><summary>Agregar o actualizar la lista</summary>
        <div class="card" style="margin-top:10px">
          <form method="post" action="/course/{course_id}/roster" enctype="multipart/form-data">
            {ROSTER_HELP}
            <textarea name="roster_text" aria-label="Lista pegada"></textarea>
            <label for="rf">O sube el archivo</label>
            <input id="rf" name="roster_file" type="file" accept=".xlsx,.csv,.txt">
            <p style="margin-top:14px"><button class="btn" type="submit">Guardar lista</button></p>
          </form>
        </div>
      </details>
      {groups_html}
      <h2 id="historial">Historial</h2>
      <div class="table-wrap"><table>
        <tr><th>Sesión</th><th>Inicio</th><th>Fin</th><th>Presentes</th><th></th></tr>
        {session_rows or "<tr><td colspan=5 class=muted>Todavía no hay sesiones.</td></tr>"}
      </table></div>
      <details style="margin-top:40px"><summary>Cambiar el nombre de la clase</summary>
        <form method="post" action="/course/{course_id}/rename" class="card" style="margin-top:10px;max-width:460px">
          <label for="new_name">Nombre nuevo</label>
          <input id="new_name" name="name" required maxlength="100" value="{e(course['name'])}">
          <p style="margin-top:12px"><button class="btn small" type="submit">Guardar nombre</button></p>
        </form>
      </details>
      <p style="margin-top:20px"><a href="/course/{course_id}/delete" style="color:var(--bad)">Borrar esta clase</a></p>""", teacher)


@app.post("/course/{course_id}/roster")
async def update_roster(course_id: int, request: Request, roster_text: str = Form(""),
                        roster_file: UploadFile | None = File(None)):
    rows, problems = await read_roster_input(roster_text, roster_file)
    with db() as conn:
        teacher = require_teacher(request, conn)
        course = own_course(conn, teacher, course_id)
        if len(_sections(rows)) > 1:
            return section_picker(teacher, f"/course/{course_id}/roster/groups",
                                  course["name"], rows, problems)
        added = save_roster(conn, course_id, rows) if rows else 0
        waiting = pending_count(conn, course_id)
    note = f"Lista guardada: {added} estudiantes." if rows else "No encontré estudiantes en lo que enviaste."
    if problems:
        note += " Revisa: " + " ".join(problems[:5])
    if added and waiting:
        return RedirectResponse(f"/course/{course_id}/validar?msg={note}", status_code=303)
    return RedirectResponse(f"/course/{course_id}?msg={note}", status_code=303)


@app.post("/course/{course_id}/roster/groups")
def update_roster_groups(course_id: int, request: Request, payload: str = Form(...),
                         sections: list[str] = Form(default=[]), name: str = Form("")):
    rows = _picked(payload, sections)
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        added = save_roster(conn, course_id, rows) if rows else 0
        waiting = pending_count(conn, course_id)
    note = f"Lista guardada: {added} estudiantes." if added else "No marcaste ningún grupo."
    if added and waiting:
        return RedirectResponse(f"/course/{course_id}/validar?msg={note}", status_code=303)
    return RedirectResponse(f"/course/{course_id}?msg={note}", status_code=303)


# ---------------------------------------------------------------
# VALIDATING AN ATTENDANCE TAKEN WITHOUT A LIST
#
# A professor without the list can still take attendance: students type a
# code and their name and wait as "pendiente". When the list arrives later,
# each pending student is matched to a row of the list, and the professor
# confirms. Found on 21 October 2026: students typed their 9-digit code or
# their username, while the Brightspace list only had an internal number.
# ---------------------------------------------------------------

STOP_WORDS = {"de", "del", "la", "las", "los", "y"}


def name_tokens(name: str) -> set:
    return set(re.findall(r"[a-z]+", _plain(name))) - STOP_WORDS


def pending_count(conn, course_id) -> int:
    return conn.execute("SELECT COUNT(*) AS n FROM enrollments WHERE course_id = %s "
                        "AND status = 'pendiente'", (course_id,)).fetchone()["n"]


def suggest_matches(conn, course_id):
    """
    For each pending student, the list row that is most likely the same
    person, and how sure:
      exacto   the typed code or username is that row's carne or username
      nombre   every word of one name is in the other, and only one row fits
      parcial  shares at least two words and no other row shares as many
    A name alone never validates anyone: the professor confirms every match.
    """
    rows = conn.execute(
        "SELECT st.*, en.status, en.section FROM students st JOIN enrollments en "
        "ON en.student_id = st.id WHERE en.course_id = %s", (course_id,)).fetchall()
    # When the list has groups, everyone who came from the list has one. An
    # approved student with no group joined by checking in, and was approved
    # by hand instead of validated (October 26: a username approved with
    # "Aprobar" left the same person twice, once without a group).
    grouped = any(r["section"] for r in rows if r["status"] == "aprobado")
    joined = (lambda r: r["status"] == "pendiente" or
              (grouped and r["status"] == "aprobado" and not r["section"]))
    listed = [r for r in rows if r["status"] == "aprobado" and not joined(r)]
    pending = sorted((r for r in rows if joined(r)), key=lambda r: r["name"])
    taken, out = set(), []
    for p in pending:
        key = normalize_carne(p["carne"])
        match, how = None, ""
        exact = [l for l in listed if key in (l["carne"], l["username"])]
        if len(exact) == 1:
            match, how = exact[0], "exacto"
        else:
            t = name_tokens(p["name"])
            subset = [l for l in listed if t and (t <= name_tokens(l["name"]) or name_tokens(l["name"]) <= t)]
            if len(subset) == 1:
                match, how = subset[0], "nombre"
            elif t:
                scored = sorted(((len(t & name_tokens(l["name"])), l) for l in listed),
                                key=lambda x: -x[0])
                if scored and scored[0][0] >= 2 and (len(scored) == 1 or scored[1][0] < scored[0][0]):
                    match, how = scored[0][1], "parcial"
        if match and match["id"] in taken:
            match, how = None, ""
        if match:
            taken.add(match["id"])
        if match or p["status"] == "pendiente":
            # An approved student with no match is simply someone new: not shown.
            out.append((p, match, how))
    return out


@app.get("/course/{course_id}/validar", response_class=HTMLResponse)
def validate_page(course_id: int, request: Request, msg: str = ""):
    with db() as conn:
        teacher = require_teacher(request, conn)
        course = own_course(conn, teacher, course_id)
        pairs = suggest_matches(conn, course_id)
    label = {"exacto": ("ok", "Mismo código o usuario"), "nombre": ("ok", "Mismo nombre"),
             "parcial": ("warn", "Nombre parecido: revisa")}
    body = ""
    for p, m, how in pairs:
        if m:
            cls, text = label[how]
            body += (f"<tr><td><input type='checkbox' name='pair' value='{p['id']}:{m['id']}' "
                     f"{'checked' if how != 'parcial' else ''} style='width:auto'></td>"
                     f"<td>{e(p['name'])}<br><span class='muted'>{e(p['carne'])}</span></td>"
                     f"<td>{e(m['name'])}<br><span class='muted'>{e(short_section(m['section']))}</span></td>"
                     f"<td><span class='pill {cls}'>{text}</span></td></tr>")
        else:
            body += (f"<tr><td></td><td>{e(p['name'])}<br><span class='muted'>{e(p['carne'])}</span></td>"
                     f"<td class='muted'>No está en la lista</td><td><span class='pill bad'>Sigue pendiente</span></td></tr>")
    flash = f"<div class='msg ok'>{e(msg)}</div>" if msg else ""
    found = sum(1 for _, m, _ in pairs if m)
    return teacher_page("Validar con la lista", f"""
      <h1>Validar “{e(course['name'])}” con la lista</h1>{flash}
      <p>{len(pairs)} estudiantes marcaron asistencia sin estar en la lista; {found} tienen
      una pareja en la lista. Revisa cada fila: al confirmar, su asistencia pasa al estudiante
      de la lista, con el código que escribió y su celular.</p>
      <form method="post" action="/course/{course_id}/validar">
        <div class="table-wrap"><table>
          <tr><th></th><th>Escribió al marcar</th><th>En la lista</th><th>Coincidencia</th></tr>
          {body or "<tr><td colspan=4 class=muted>No hay pendientes.</td></tr>"}
        </table></div>
        <p style="margin-top:16px;display:flex;gap:8px;flex-wrap:wrap">
          <button class="btn" type="submit">Confirmar los marcados</button>
          <a class="btn grey" href="/course/{course_id}">Volver</a></p>
      </form>""", teacher)


@app.post("/course/{course_id}/validar")
def validate_confirm(course_id: int, request: Request, pair: list[str] = Form(default=[])):
    """
    Merge each confirmed pending student into their list row. The list's name
    wins; the code the student typed is kept as their carne when it is a
    number they know (the list may only carry an internal id); the username
    stays as a second way in. Nothing is guessed: only checked rows move.
    """
    done = 0
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        valid = {(p["id"], m["id"]) for p, m, _ in suggest_matches(conn, course_id) if m}
        for item in pair:
            try:
                pid, lid = (int(x) for x in item.split(":"))
            except ValueError:
                continue
            if (pid, lid) not in valid:
                continue
            p = conn.execute("SELECT * FROM students WHERE id = %s", (pid,)).fetchone()
            l = conn.execute("SELECT * FROM students WHERE id = %s", (lid,)).fetchone()
            typed = p["carne"]
            merge_students(conn, wrong=p, right=l)
            if typed.isdigit() and typed != l["carne"]:
                conn.execute("UPDATE students SET carne = %s WHERE id = %s", (typed, lid))
            elif not typed.isdigit() and not l["username"]:
                # The username the student typed stays as a second way in, so
                # next time it finds them directly.
                conn.execute("UPDATE students SET username = %s WHERE id = %s "
                             "AND NOT EXISTS (SELECT 1 FROM students WHERE username = %s)",
                             (typed, lid, typed))
            log_event(conn, None, lid, "validado_con_lista",
                      f"{typed} → {l['name']}, por {teacher['name']}")
            done += 1
    return RedirectResponse(f"/course/{course_id}?msg={done} estudiantes validados con la lista.",
                            status_code=303)


@app.get("/course/{course_id}/delete", response_class=HTMLResponse)
def delete_course_page(course_id: int, request: Request, error: str = ""):
    with db() as conn:
        teacher = require_teacher(request, conn)
        course = own_course(conn, teacher, course_id)
        counts = conn.execute("""
            SELECT (SELECT COUNT(*) FROM enrollments WHERE course_id = %s) AS students,
                   (SELECT COUNT(*) FROM sessions WHERE course_id = %s) AS sessions
        """, (course_id, course_id)).fetchone()
    msg = "<div class='msg bad'>El nombre no coincide. No se borró nada.</div>" if error else ""
    return teacher_page("Borrar clase", f"""
      <h1>Borrar “{e(course['name'])}”</h1>{msg}
      <div class="card" style="max-width:560px">
        <p>Se borran la lista ({counts['students']} estudiantes), las {counts['sessions']} sesiones
        y toda su asistencia. <strong>No se puede deshacer.</strong> Si necesitas los reportes,
        descárgalos antes.</p>
        <form method="post" action="/course/{course_id}/delete">
          <label for="confirm_name">Escribe el nombre de la clase para confirmar</label>
          <input id="confirm_name" name="confirm_name" required autocomplete="off">
          <p style="margin-top:16px;display:flex;gap:8px;flex-wrap:wrap">
            <button class="btn red" type="submit">Borrar para siempre</button>
            <a class="btn grey" href="/course/{course_id}">Cancelar</a></p>
        </form>
      </div>""", teacher)


@app.post("/course/{course_id}/delete")
def delete_course(course_id: int, request: Request, confirm_name: str = Form(...)):
    """
    Typing the class name is the confirmation: a misplaced tap cannot erase a
    semester of attendance.

    Students are shared between classes, so they are not deleted with the
    class. Only the ones left with no class and no attendance anywhere go,
    which removes test carnes without touching a student someone else teaches.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        course = own_course(conn, teacher, course_id)
        if confirm_name.strip().casefold() != course["name"].strip().casefold():
            return RedirectResponse(f"/course/{course_id}/delete?error=1", status_code=303)
        conn.execute("DELETE FROM courses WHERE id = %s", (course_id,))
        conn.execute("""
            DELETE FROM students st
            WHERE NOT EXISTS (SELECT 1 FROM enrollments en WHERE en.student_id = st.id)
              AND NOT EXISTS (SELECT 1 FROM attendance a WHERE a.student_id = st.id)
        """)
    return RedirectResponse(f"/?msg=Clase “{course['name']}” borrada.", status_code=303)


def class_student(conn, course_id, student_id):
    return conn.execute(
        "SELECT st.* FROM students st JOIN enrollments en ON en.student_id = st.id "
        "WHERE st.id = %s AND en.course_id = %s", (student_id, course_id)).fetchone()


@app.get("/course/{course_id}/student/{student_id}/edit", response_class=HTMLResponse)
def edit_student_page(course_id: int, student_id: int, request: Request, error: str = ""):
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        st = class_student(conn, course_id, student_id)
        if not st:
            return RedirectResponse(f"/course/{course_id}", status_code=303)
    msg = f"<div class='msg bad'>{e(error)}</div>" if error else ""
    return teacher_page("Corregir estudiante", f"""
      <h1>Corregir estudiante</h1>{msg}
      <div class="card" style="max-width:520px">
        <form method="post" action="/course/{course_id}/student/{student_id}/edit">
          <label for="carne">Carné</label>
          <input id="carne" name="carne" value="{e(st['carne'])}" required autocapitalize="characters">
          <label for="name">Nombre</label>
          <input id="name" name="name" value="{e(st['name'])}" required maxlength="80">
          <p class="hint">Su celular sigue vinculado y su asistencia se conserva. Si el carné
          correcto ya existe (por ejemplo, porque está en la lista), los dos registros se unen.</p>
          <p style="margin-top:16px;display:flex;gap:8px;flex-wrap:wrap">
            <button class="btn" type="submit">Guardar</button>
            <a class="btn grey" href="/course/{course_id}">Cancelar</a></p>
        </form>
      </div>""", teacher)


@app.post("/course/{course_id}/student/{student_id}/edit")
def edit_student(course_id: int, student_id: int, request: Request,
                 carne: str = Form(...), name: str = Form(...)):
    """
    A student who mistypes the carne in a class without a list cannot fix it
    alone: from then on the browser marks for the wrong carne and ignores what
    is typed. The professor corrects it here.

    If the right carne already belongs to another record (the student is on
    the list, or linked properly in another class), the two are merged:
    attendance, classes and history move to the right record, and the
    browser follows unless that record already has its own.
    """
    new_carne = normalize_carne(carne)
    new_name = re.sub(r"\s+", " ", name).strip()[:80]
    back = f"/course/{course_id}/student/{student_id}/edit"
    if not new_carne or not new_name:
        return RedirectResponse(f"{back}?error=Escribe el carné y el nombre.", status_code=303)
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        st = class_student(conn, course_id, student_id)
        if not st:
            return RedirectResponse(f"/course/{course_id}", status_code=303)
        target = conn.execute("SELECT * FROM students WHERE carne = %s AND id <> %s",
                              (new_carne, student_id)).fetchone()
        if target is None:
            conn.execute("UPDATE students SET carne = %s, name = %s WHERE id = %s",
                         (new_carne, new_name, student_id))
            keep, note = student_id, f"Datos de {new_name} corregidos."
        else:
            merge_students(conn, wrong=st, right=target)
            keep, note = target["id"], f"{st['carne']} unido a {target['carne']} ({target['name']})."
        log_event(conn, None, keep, "carne_corregido",
                  f"{st['carne']} → {new_carne}, por {teacher['name']}")
    return RedirectResponse(f"/course/{course_id}?msg={note}", status_code=303)


def merge_students(conn, wrong, right):
    """Move everything from the mistyped record to the right one, then drop it."""
    w, r = wrong["id"], right["id"]
    # Attendance: keep the right record's row where both marked the same class.
    conn.execute("DELETE FROM attendance a WHERE a.student_id = %s AND EXISTS ("
                 "SELECT 1 FROM attendance b WHERE b.student_id = %s "
                 "AND b.session_id = a.session_id)", (w, r))
    conn.execute("UPDATE attendance SET student_id = %s WHERE student_id = %s", (r, w))
    # Classes: an approval on either side counts.
    conn.execute("UPDATE enrollments e SET status = 'aprobado' WHERE e.student_id = %s "
                 "AND EXISTS (SELECT 1 FROM enrollments f WHERE f.student_id = %s "
                 "AND f.course_id = e.course_id AND f.status = 'aprobado')", (r, w))
    conn.execute("DELETE FROM enrollments e WHERE e.student_id = %s AND EXISTS ("
                 "SELECT 1 FROM enrollments f WHERE f.student_id = %s "
                 "AND f.course_id = e.course_id)", (w, r))
    conn.execute("UPDATE enrollments SET student_id = %s WHERE student_id = %s", (r, w))
    conn.execute("UPDATE device_changes SET student_id = %s WHERE student_id = %s", (r, w))
    conn.execute("UPDATE events SET student_id = %s WHERE student_id = %s", (r, w))
    # The browser follows only if the right record has none of its own.
    device = wrong["device_id"]
    conn.execute("DELETE FROM students WHERE id = %s", (w,))
    if device and not right["device_id"]:
        conn.execute("UPDATE students SET device_id = %s, device_no = device_no + 1, "
                     "registered_at = COALESCE(registered_at, now()) WHERE id = %s", (device, r))


@app.post("/course/{course_id}/unlink")
def unlink_device(course_id: int, request: Request, student_id: int = Form(...)):
    """
    For a stolen or replaced phone, or a carne linked by mistake during a test.
    The student types the carne again on the next scan, and that first link
    raises no "changed device" alert, because there is no device to change from.
    Only a professor whose class includes the student can do it.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        student = conn.execute(
            "SELECT st.* FROM students st JOIN enrollments en ON en.student_id = st.id "
            "WHERE st.id = %s AND en.course_id = %s", (student_id, course_id)).fetchone()
        if not student:
            return RedirectResponse(f"/course/{course_id}", status_code=303)
        conn.execute("UPDATE students SET device_id = NULL WHERE id = %s", (student_id,))
        log_event(conn, None, student_id, "desvinculado", f"por {teacher['name']}")
    return RedirectResponse(f"/course/{course_id}?msg=Celular de {student['name']} desvinculado.",
                            status_code=303)


@app.post("/course/{course_id}/approve")
def approve(course_id: int, request: Request, student_id: int = Form(0), all: str = Form("")):
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        if all:
            conn.execute("UPDATE enrollments SET status = 'aprobado' WHERE course_id = %s",
                         (course_id,))
        else:
            conn.execute("UPDATE enrollments SET status = 'aprobado' "
                         "WHERE course_id = %s AND student_id = %s", (course_id, student_id))
    return RedirectResponse(f"/course/{course_id}", status_code=303)


@app.post("/course/{course_id}/remove")
def remove(course_id: int, request: Request, student_id: int = Form(...)):
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        conn.execute("DELETE FROM enrollments WHERE course_id = %s AND student_id = %s",
                     (course_id, student_id))
        # Data that serves no class is not kept: a person removed from their
        # only class, with no attendance left, is deleted entirely.
        if not conn.execute("SELECT 1 FROM enrollments WHERE student_id = %s UNION ALL "
                            "SELECT 1 FROM attendance WHERE student_id = %s LIMIT 1",
                            (student_id, student_id)).fetchone():
            conn.execute("DELETE FROM students WHERE id = %s", (student_id,))
    return RedirectResponse(f"/course/{course_id}", status_code=303)


# ---------------------------------------------------------------
# READING A CLASS LIST
#
# Same idea as generate_codes.py: accept whatever the professor has. A paste
# from Excel arrives separated by tabs; a CSV exported from Excel in Spanish
# uses semicolons.
# ---------------------------------------------------------------

def _plain(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return text.strip().lower()


# Column names seen in the wild. Banner, the university's academic system, exports
# "SPRIDEN_ID, SPRIDEN_LAST_NAME, SSBSECT_CRN": the student code, the full
# name and the group (CRN). Brightspace (its course platform) exports
# "OrgDefinedId, Last Name, First Name, Email, Sections"; Spanish sheets use
# "Código, Apellidos, Nombres"; a paste from anywhere may have no header at all.
FULL_NAME_HEADERS = {"name", "full name", "estudiante", "alumno", "nombre completo",
                     "apellidos y nombres", "nombres y apellidos", "spriden_last_name"}
FIRST_NAME_HEADERS = {"first name", "nombre", "nombres", "primer nombre"}
LAST_NAME_HEADERS = {"last name", "apellido", "apellidos"}
CARNE_HEADERS = {"carne", "carnet", "codigo", "codigo estudiantil", "codigo de estudiante",
                 "id", "student id", "id estudiante", "orgdefinedid", "org defined id",
                 "matricula", "spriden_id"}
USERNAME_HEADERS = {"username", "usuario", "user name"}
EMAIL_HEADERS = {"email", "correo", "correo electronico", "e-mail", "email address"}
# "documento" and "cedula" were removed on purpose: a national ID number is
# sensitive personal data and is never a good key for attendance.
SECTION_HEADERS = {"sections", "section", "seccion", "secciones", "grupo", "nrc", "crn",
                   "ssbsect_crn"}
# From an email only the username before the @ is kept; the address itself
# is never stored, and data the system does not keep cannot leak.


def _looks_like_carne(cell: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9\-]*\d{4,}[A-Za-z0-9\-]*", cell))


def _tidy_name(name: str) -> str:
    """Platforms export names in capitals; the report reads better in title case."""
    name = re.sub(r"\s+", " ", name).strip()
    if name and name == name.upper():
        name = name.title()
    return name[:80]


def parse_roster_rows(rows):
    """
    Returns ([(carne, name, section, username), ...], [problems]).

    With headers, the columns are found by name; first and last name in
    separate columns are joined. Without headers, the cell that looks like an
    ID (four or more digits, no spaces) is the carne and the rest is the name.
    `section` is the group the row belongs to, or "" when the file has none.
    """
    rows = [[normalize_cell(c) for c in r] for r in rows if any(normalize_cell(c) for c in r)]
    if not rows:
        return [], []
    header = [_plain(c) for c in rows[0]]

    def col(names):
        return next((i for i, h in enumerate(header) if h in names), None)

    full_col, first_col, last_col = col(FULL_NAME_HEADERS), col(FIRST_NAME_HEADERS), col(LAST_NAME_HEADERS)
    carne_col, section_col = col(CARNE_HEADERS), col(SECTION_HEADERS)
    user_col, email_col = col(USERNAME_HEADERS), col(EMAIL_HEADERS)

    result, problems, seen = [], [], set()

    def cell(r, i):
        return r[i] if i is not None and i < len(r) else ""

    def add(carne, name, section, line, username="", last=""):
        username = normalize_carne(username)
        carne = normalize_carne(carne) or username
        name = _tidy_name(name)
        if not carne:
            problems.append(f"fila {line}: sin carné.")
            return
        if not name:
            problems.append(f"fila {line}: sin nombre.")
            return
        if carne in seen:
            problems.append(f"fila {line}: carné {carne} repetido.")
            return
        seen.add(carne)
        result.append((carne, name, section.strip(), username, _tidy_name(last) if last else ""))

    has_name = full_col is not None or first_col is not None or last_col is not None
    has_key = carne_col is not None or user_col is not None or email_col is not None
    if has_name or has_key:
        if not has_name or not has_key:
            missing = "nombre" if not has_name else "carné o código"
            return [], [f"La lista tiene encabezados pero falta la columna de {missing}."]
        for n, r in enumerate(rows[1:], start=2):
            if full_col is not None:
                name = cell(r, full_col)
            else:
                name = f"{cell(r, first_col)} {cell(r, last_col)}"
            username = cell(r, user_col) or cell(r, email_col)
            add(cell(r, carne_col), name, cell(r, section_col), n, username, cell(r, last_col))
        return result, problems

    for n, r in enumerate(rows, start=1):
        cells = [c for c in r if c]
        carne = next((c for c in cells if _looks_like_carne(c)), "")
        name = " ".join(c for c in cells if c != carne and "@" not in c)
        add(carne, name, "", n)
    return result, problems


def normalize_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


async def read_roster_input(text: str, upload: UploadFile | None):
    if upload is not None and upload.filename:
        data = await upload.read()
        if upload.filename.lower().endswith(".xlsx"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
                return parse_roster_rows(rows)
            except Exception:
                return [], ["No pude leer ese archivo de Excel."]
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
    text = (text or "").strip()
    if not text:
        return [], []
    if "\t" in text:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;").delimiter
        except csv.Error:
            delimiter = ","
    return parse_roster_rows(list(csv.reader(io.StringIO(text), delimiter=delimiter)))


def save_roster(conn, course_id, rows) -> int:
    """
    The list from the university is authoritative: if a student typed their
    own name differently when checking in, the list's spelling wins, and
    anyone on the list who was waiting for approval becomes approved.
    """
    listed_codes = {r[0] for r in rows}
    in_class = conn.execute(
        "SELECT st.* FROM students st JOIN enrollments en ON en.student_id = st.id "
        "WHERE en.course_id = %s AND en.status = 'aprobado'", (course_id,)).fetchall()
    for carne, name, *rest in rows:
        section = rest[0] if rest else ""
        username = (rest[1] if len(rest) > 1 else "") or None
        last = (rest[2] if len(rest) > 2 else "") or None
        # A student validated earlier may now carry the code they typed instead
        # of the list's internal number; the username still finds them.
        existing = conn.execute(
            "SELECT id FROM students WHERE carne = %s OR (%s::text IS NOT NULL AND username = %s) "
            "ORDER BY (carne = %s) DESC LIMIT 1", (carne, username, username, carne)).fetchone()
        if not existing:
            # A second list for the same class may use another code: Brightspace
            # gives an internal 6-digit id, Banner the 9-digit student code.
            # The same person is found by name, among students already listed
            # in this class (never a pending one: those the professor confirms),
            # when exactly one fits and their current code is not on this list.
            # Their attendance stays; the code becomes the new one.
            t = name_tokens(name)
            same = [st for st in in_class if st["carne"] not in listed_codes and t and
                    (t <= name_tokens(st["name"]) or name_tokens(st["name"]) <= t)]
            if len(same) == 1:
                existing = same[0]
                log_event(conn, None, existing["id"], "carne_corregido",
                          f"{existing['carne']} → {carne} (lista nueva)")
                conn.execute("UPDATE students SET carne = %s WHERE id = %s", (carne, existing["id"]))
                in_class = [st for st in in_class if st["id"] != existing["id"]]
        if existing:
            student = conn.execute(
                "UPDATE students SET name = %s, username = COALESCE(%s, username), "
                "last_name = COALESCE(%s, last_name) WHERE id = %s RETURNING id",
                (name, username, last, existing["id"])).fetchone()
        else:
            student = conn.execute(
                "INSERT INTO students (carne, name, username, last_name) VALUES (%s, %s, %s, %s) "
                "RETURNING id", (carne, name, username, last)).fetchone()
        conn.execute(
            "INSERT INTO enrollments (course_id, student_id, status, section) "
            "VALUES (%s, %s, 'aprobado', %s) "
            "ON CONFLICT (course_id, student_id) DO UPDATE SET status = 'aprobado', "
            "section = EXCLUDED.section",
            (course_id, student["id"], section),
        )
    conn.execute("UPDATE courses SET has_roster = TRUE WHERE id = %s", (course_id,))
    return len(rows)


# ---------------------------------------------------------------
# SESSIONS AND THE LIVE SCREEN
# ---------------------------------------------------------------

def course_sections(conn, course_id):
    """Groups in a class with their sizes, e.g. {'SISTEMA NERVIOSO2026303614': 26}."""
    return {r["section"]: r["n"] for r in conn.execute(
        "SELECT section, COUNT(*) AS n FROM enrollments "
        "WHERE course_id = %s AND section <> '' GROUP BY section ORDER BY section",
        (course_id,)).fetchall()}


@app.post("/course/{course_id}/start")
def start_session(course_id: int, request: Request, sections: list[str] = Form(default=[]),
                  chosen: str = Form("")):
    """
    One class, several groups: a professor may teach groups 3614-3616 at 7:00
    and 3617-3619 at 9:00 from the same list. When the class has more than one
    group, she picks today's groups first, and the absent list counts only
    those. The choice is saved with the session, so the report knows it too.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        course = own_course(conn, teacher, course_id)
        groups = course_sections(conn, course_id)
        if len(groups) > 1 and not chosen:
            boxes = "".join(
                f"<label style='font-weight:normal;display:flex;gap:6px;align-items:center;margin:0;"
                f"padding:6px 10px;border:1px solid var(--line);border-radius:8px;background:#fff'>"
                f"<input type='checkbox' name='sections' value='{e(sec)}' style='width:auto'>"
                f"{e(short_section(sec))} <span class='muted'>· {n}</span></label>"
                for sec, n in groups.items())
            return HTMLResponse(teacher_page("Grupos de hoy", f"""
              <h1>{e(course['name'])}</h1>
              <p style="font-size:18px">¿Qué grupos están en el salón ahora?</p>
              <form method="post" action="/course/{course_id}/start" style="max-width:520px">
                <input type="hidden" name="chosen" value="1">
                <p class="hint" style="margin-top:0">Si no marcas ninguno, se cuentan todos los grupos.</p>
                <div style="display:flex;flex-wrap:wrap;gap:8px">{boxes}</div>
                <p style="margin-top:14px"><button class="btn big" type="submit">Tomar asistencia</button></p>
              </form>""", teacher))
        picked = [sec for sec in sections if sec in groups] or None
        # A professor teaches one class at a time. Anything left open from an
        # earlier class is closed, so it cannot keep accepting check-ins.
        conn.execute("""
            UPDATE sessions SET active = FALSE, ended_at = now()
            WHERE active AND course_id IN (SELECT id FROM courses WHERE teacher_id = %s)
        """, (teacher["id"],))
        session = conn.execute(
            "INSERT INTO sessions (course_id, sections) VALUES (%s, %s) RETURNING id",
            (course_id, picked)).fetchone()
        conn.execute("UPDATE courses SET last_used_at = now() WHERE id = %s", (course_id,))
    return RedirectResponse(f"/session/{session['id']}/live", status_code=303)


LIVE_PAGE = """
<style>
  .live-grid { display:grid; grid-template-columns:minmax(0,1.1fr) minmax(0,1fr); gap:20px; align-items:start; }
  @media (max-width:820px) { .live-grid { grid-template-columns:1fr; } }
  #qr { width:100%; max-width:440px; aspect-ratio:1; display:block; margin:0 auto;
        background:#fff; border:1px solid var(--line); border-radius:10px; }
  #count { font-size:64px; font-weight:800; line-height:1; font-variant-numeric:tabular-nums; }
  .list { display:grid; gap:6px; margin-top:8px; max-height:260px; overflow:auto; }
  .absent-btn { text-align:left; background:#fff; border:1px solid var(--line); border-radius:8px;
                padding:8px 10px; font:inherit; cursor:pointer; }
  .absent-btn:hover { border-color:var(--ok); }
  .alert { background:var(--warn-soft); border-radius:8px; padding:8px 10px; font-size:14px; }
  .recent { font-size:14px; }
  #details[hidden] { display:none !important; }
</style>
<div style="display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap">
  <h1>__COURSE__</h1>
  <form method="post" action="/session/__SID__/end"
        onsubmit="return confirm('¿Terminar la asistencia de esta clase?')">
    <button class="btn red" type="submit">Terminar asistencia</button></form>
</div>
<div class="live-grid">
  <div class="card" style="text-align:center">
    <p style="margin:0 0 8px;font-size:20px;font-weight:700">Escanea con la cámara para marcar asistencia</p>
    <img id="qr" alt="Código QR de asistencia">
    <p class="hint">Cambia cada __ROT__ segundos. No hace falta instalar nada.</p>
  </div>
  <div style="display:grid;gap:14px">
    <div class="card">
      <div class="muted">Presentes</div>
      <div id="count">0</div>
      <div id="pending" class="hint"></div>
    </div>
    <!-- The live screen is usually projected, so names stay hidden until the
         professor opens them. Students in the room only see the QR and the count. -->
    <div class="card">
      <button id="toggle" class="btn light" type="button" onclick="toggleDetails()">Mostrar ausentes y alertas</button>
      <p class="hint" style="margin:6px 0 0">Muestra nombres. Ábrelo cuando no estés proyectando.</p>
    </div>
    <div id="details" style="display:grid;gap:14px" hidden>
    <div class="card">
      <strong>Por revisar</strong>
      <div id="alerts" class="list"><span class="muted">Nada por ahora.</span></div>
    </div>
    <div class="card">
      <strong>Ausentes</strong> <span class="hint">· toca un nombre para marcarlo a mano</span>
      <div id="absent" class="list"></div>
      <details style="margin-top:10px"><summary>Agregar a alguien que no está en la lista</summary>
        <form id="addForm" style="margin-top:8px">
          <input id="addCarne" placeholder="Carné" required>
          <input id="addName" placeholder="Nombre" style="margin-top:6px">
          <button class="btn small" type="submit" style="margin-top:6px">Agregar y marcar</button>
        </form>
      </details>
      <p id="manualMsg" class="hint"></p>
    </div>
    <div class="card recent"><strong>Últimos en marcar</strong><div id="recent" class="list"></div></div>
    </div>
  </div>
</div>
<script>
const SID = __SID__, ROT = __ROT__;
function refreshQR() { document.getElementById('qr').src = '/session/' + SID + '/qr?t=' + Date.now(); }
refreshQR(); setInterval(refreshQR, ROT * 1000);

let open = false, alertCount = 0;
function toggleDetails() {
  open = !open;
  document.getElementById('details').hidden = !open;
  label();
}
function label() {
  const n = alertCount ? ' (' + alertCount + ' por revisar)' : '';
  document.getElementById('toggle').textContent = (open ? 'Ocultar' : 'Mostrar') + ' ausentes y alertas' + n;
}

function el(tag, text, cls) { const n = document.createElement(tag); n.textContent = text; if (cls) n.className = cls; return n; }

async function poll() {
  const r = await fetch('/session/' + SID + '/state');
  if (r.status === 401) { location = '/login'; return; }
  if (!r.ok) return;
  const d = await r.json();
  if (!d.active) { location = '/session/' + SID + '/report'; return; }
  document.getElementById('count').textContent = d.present + ' de ' + d.total;
  document.getElementById('pending').textContent =
    d.pending ? d.pending + ' estudiante(s) nuevos por aprobar al final' : '';

  alertCount = d.alerts.filter(a => !a.decision).length; label();
  const alerts = document.getElementById('alerts'); alerts.innerHTML = '';
  if (!d.alerts.length) alerts.appendChild(el('span', 'Nada por ahora.', 'muted'));
  d.alerts.forEach(a => {
    const box = el('div', a.time + ' · ' + (a.name ? a.name + ' · ' : '') + a.label + (a.detail ? ' (' + a.detail + ')' : ''), 'alert');
    if (a.decision) {
      box.appendChild(el('strong', ' · ' + a.decision));
    } else if (a.student) {
      [['confirmado', 'Confirmar'], ['rechazado', 'Rechazar']].forEach(([value, text]) => {
        const b = el('button', text, 'btn small ' + (value === 'rechazado' ? 'red' : 'light'));
        b.type = 'button'; b.style.marginLeft = '6px';
        b.onclick = () => decide(a.id, value);
        box.appendChild(b);
      });
    }
    alerts.appendChild(box);
  });

  const absent = document.getElementById('absent'); absent.innerHTML = '';
  if (!d.absent.length) absent.appendChild(el('span', d.total ? 'Nadie. Están todos.' : 'Sin lista: aparecerán al marcar.', 'muted'));
  d.absent.forEach(s => {
    const b = el('button', s.name + ' · ' + s.carne, 'absent-btn');
    b.type = 'button';
    b.onclick = () => manual({ student_id: s.id }, s.name);
    absent.appendChild(b);
  });

  const recent = document.getElementById('recent'); recent.innerHTML = '';
  d.recent.forEach(s => recent.appendChild(el('div', s.time + ' · ' + s.name + (s.method === 'manual' ? ' (a mano)' : ''))));
}

async function decide(eventId, decision) {
  const f = new FormData(); f.append('decision', decision);
  await fetch('/session/' + SID + '/event/' + eventId + '/decide', { method: 'POST', body: f });
  poll();
}

async function manual(fields, label) {
  const f = new FormData();
  Object.entries(fields).forEach(([k, v]) => f.append(k, v));
  const r = await fetch('/session/' + SID + '/manual', { method: 'POST', body: f });
  const d = await r.json().catch(() => ({ message: 'No se pudo marcar.' }));
  document.getElementById('manualMsg').textContent = d.message;
  poll();
}

document.getElementById('addForm').addEventListener('submit', ev => {
  ev.preventDefault();
  manual({ carne: document.getElementById('addCarne').value, name: document.getElementById('addName').value });
  ev.target.reset();
});

poll(); setInterval(poll, 3000);
</script>
"""


@app.get("/session/{session_id}/live", response_class=HTMLResponse)
def live_page(session_id: int, request: Request):
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
    if not session["active"]:
        return RedirectResponse(f"/session/{session_id}/report", status_code=303)
    title = session["course_name"]
    if session["sections"]:
        title += " · " + ", ".join(short_section(g) for g in session["sections"])
    body = (LIVE_PAGE.replace("__SID__", str(session_id))
            .replace("__ROT__", str(QR_ROTATE_SECONDS))
            .replace("__COURSE__", e(title)))
    return teacher_page(session["course_name"], body, teacher)


@app.get("/session/{session_id}/state")
def live_state(session_id: int, request: Request):
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        # Only the groups chosen for this session count toward "X de Y".
        groups = session["sections"]
        total = conn.execute(
            "SELECT COUNT(*) AS n FROM enrollments WHERE course_id = %s "
            "AND (%s::text[] IS NULL OR section = ANY(%s::text[]))",
            (session["course_id"], groups, groups)).fetchone()["n"]
        pending = conn.execute("SELECT COUNT(*) AS n FROM enrollments WHERE course_id = %s "
                               "AND status = 'pendiente'", (session["course_id"],)).fetchone()["n"]
        present = conn.execute("SELECT COUNT(*) AS n FROM attendance WHERE session_id = %s",
                               (session_id,)).fetchone()["n"]
        absent = conn.execute("""
            SELECT st.id, st.name, st.carne FROM enrollments en
            JOIN students st ON st.id = en.student_id
            WHERE en.course_id = %s
              AND (%s::text[] IS NULL OR en.section = ANY(%s::text[]))
              AND NOT EXISTS (
                SELECT 1 FROM attendance a WHERE a.session_id = %s AND a.student_id = st.id)
            ORDER BY st.name
        """, (session["course_id"], groups, groups, session_id)).fetchall()
        recent = conn.execute("""
            SELECT st.name, a.marked_at, a.method FROM attendance a
            JOIN students st ON st.id = a.student_id
            WHERE a.session_id = %s ORDER BY a.marked_at DESC LIMIT 8
        """, (session_id,)).fetchall()
        alerts = conn.execute("""
            SELECT ev.*, st.name FROM events ev LEFT JOIN students st ON st.id = ev.student_id
            WHERE ev.session_id = %s AND ev.review
            ORDER BY (ev.kind = 'rafaga'), ev.created_at DESC LIMIT 30
        """, (session_id,)).fetchall()
    return {
        "active": session["active"],
        "present": present, "total": total, "pending": pending,
        "absent": [{"id": a["id"], "name": a["name"], "carne": a["carne"]} for a in absent],
        "recent": [{"name": r["name"], "time": local_time(r["marked_at"]), "method": r["method"]}
                   for r in recent],
        "alerts": [{"id": a["id"], "time": local_time(a["created_at"]), "name": a["name"] or "",
                    "label": EVENT_LABELS.get(a["kind"], a["kind"]), "detail": a["detail"],
                    "student": a["student_id"] is not None, "decision": a["decision"] or ""}
                   for a in alerts],
    }


@app.post("/session/{session_id}/manual")
def manual_mark(session_id: int, request: Request, student_id: int = Form(0),
                carne: str = Form(""), name: str = Form("")):
    """
    For the dead battery, the phone with no data, or the student not on the
    list. Recorded as manual and as an event, so the report says who was
    marked by hand.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        # A professor can still mark someone by hand after ending the class
        # (from the report). Students cannot: the QR stops working.
        if student_id:
            student = conn.execute(
                "SELECT st.* FROM students st JOIN enrollments en ON en.student_id = st.id "
                "WHERE st.id = %s AND en.course_id = %s", (student_id, session["course_id"]),
            ).fetchone()
            if not student:
                return {"message": "Ese estudiante no está en esta clase."}
        else:
            carne = normalize_carne(carne)
            if not carne:
                return {"message": "Escribe el carné."}
            student = find_student(conn, carne)
            if not student:
                if not name.strip():
                    return {"message": "Ese carné es nuevo: escribe también el nombre."}
                student = conn.execute(
                    "INSERT INTO students (carne, name) VALUES (%s, %s) RETURNING *",
                    (carne, name.strip()[:80]),
                ).fetchone()
            conn.execute(
                "INSERT INTO enrollments (course_id, student_id, status) VALUES (%s, %s, 'aprobado') "
                "ON CONFLICT (course_id, student_id) DO UPDATE SET status = 'aprobado'",
                (session["course_id"], student["id"]),
            )
        _, created = record_attendance(conn, session, student, "manual")
        if created:
            log_event(conn, session_id, student["id"], "manual", f"por {teacher['name']}")
    return {"message": f"{student['name']}: {'marcado a mano' if created else 'ya estaba marcado'}."}


@app.post("/session/{session_id}/event/{event_id}/decide")
def decide_event(session_id: int, event_id: int, request: Request, decision: str = Form(...)):
    """
    The professor checks the student in person and confirms or rejects.
    A rejected check-in stays in the database and in the report, marked as
    rejected: nothing is erased, so the decision itself can be audited.
    """
    if decision not in ("confirmado", "rechazado"):
        return JSONResponse({"error": "decision"}, status_code=400)
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_session(conn, teacher, session_id)
        conn.execute(
            "UPDATE events SET decision = %s, decided_at = now(), decided_by = %s "
            "WHERE id = %s AND session_id = %s AND review AND student_id IS NOT NULL",
            (decision, teacher["name"], event_id, session_id))
    return {"ok": True}


@app.post("/course/{course_id}/rename")
def rename_course(course_id: int, request: Request, name: str = Form(...)):
    """A class created as "Prueba" can become "Fundamentos" without losing anything."""
    new = re.sub(r"\s+", " ", name).strip()[:100]
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        if new:
            conn.execute("UPDATE courses SET name = %s WHERE id = %s", (new, course_id))
    return RedirectResponse(f"/course/{course_id}?msg=Nombre guardado.", status_code=303)


@app.get("/session/{session_id}/delete", response_class=HTMLResponse)
def delete_session_page(session_id: int, request: Request, error: str = ""):
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        n = conn.execute("SELECT COUNT(*) AS n FROM attendance WHERE session_id = %s",
                         (session_id,)).fetchone()["n"]
    msg = "<div class='msg bad'>El número no coincide. No se borró nada.</div>" if error else ""
    return teacher_page("Borrar sesión", f"""
      <h1>Borrar la sesión #{session_id} de “{e(session['course_name'])}”</h1>{msg}
      <div class="card" style="max-width:560px">
        <p>Se borran esta toma de asistencia y sus {n} registros. Los estudiantes que solo
        existían por esta sesión (marcaron sin estar en la lista y no tienen otra asistencia)
        se borran también, con sus datos. <strong>No se puede deshacer.</strong>
        Si necesitas el reporte, descárgalo antes.</p>
        <form method="post" action="/session/{session_id}/delete">
          <label for="confirm_id">Escribe el número de la sesión ({session_id}) para confirmar</label>
          <input id="confirm_id" name="confirm_id" required autocomplete="off" inputmode="numeric">
          <p style="margin-top:16px;display:flex;gap:8px;flex-wrap:wrap">
            <button class="btn red" type="submit">Borrar para siempre</button>
            <a class="btn grey" href="/course/{session['course_id']}/historial">Cancelar</a></p>
        </form>
      </div>""", teacher)


@app.post("/session/{session_id}/delete")
def delete_session(session_id: int, request: Request, confirm_id: str = Form(...)):
    """
    A test session, for example one taken with someone's national ID, can be
    removed on its own. Students who exist only because of it go with it, so
    no personal data from the test stays behind.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        course_id = session["course_id"]
        if confirm_id.strip().lstrip("#") != str(session_id):
            return RedirectResponse(f"/session/{session_id}/delete?error=1", status_code=303)
        marked = [r["student_id"] for r in conn.execute(
            "SELECT student_id FROM attendance WHERE session_id = %s", (session_id,)).fetchall()]
        conn.execute("DELETE FROM sessions WHERE id = %s", (session_id,))
        for sid in marked:
            if conn.execute("SELECT 1 FROM attendance WHERE student_id = %s", (sid,)).fetchone():
                continue
            # Pending in this class only because of the deleted session: out.
            conn.execute("DELETE FROM enrollments WHERE course_id = %s AND student_id = %s "
                         "AND status = 'pendiente'", (course_id, sid))
            if not conn.execute("SELECT 1 FROM enrollments WHERE student_id = %s", (sid,)).fetchone():
                conn.execute("DELETE FROM students WHERE id = %s", (sid,))
    return RedirectResponse(f"/course/{course_id}/historial", status_code=303)


@app.post("/session/{session_id}/groups")
def set_session_groups(session_id: int, request: Request, sections: list[str] = Form(default=[])):
    """Which groups were in the room, set after the fact. Nothing is deleted."""
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        valid = set(course_sections(conn, session["course_id"]))
        chosen = [g for g in sections if g in valid]
        conn.execute("UPDATE sessions SET sections = %s WHERE id = %s",
                     (chosen or None, session_id))
    return RedirectResponse(f"/session/{session_id}/report", status_code=303)


@app.post("/course/{course_id}/groups/remove")
def remove_group(course_id: int, request: Request, section: str = Form(...)):
    """
    A group uploaded by mistake leaves the class. Students of that group who
    already have attendance in this class stay, so no attendance is lost.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_course(conn, teacher, course_id)
        gone = conn.execute("""
            DELETE FROM enrollments en WHERE en.course_id = %s AND en.section = %s
              AND NOT EXISTS (SELECT 1 FROM attendance a JOIN sessions s ON s.id = a.session_id
                              WHERE s.course_id = %s AND a.student_id = en.student_id)
            RETURNING en.student_id""", (course_id, section, course_id)).fetchall()
        kept = conn.execute("SELECT COUNT(*) AS n FROM enrollments WHERE course_id = %s "
                            "AND section = %s", (course_id, section)).fetchone()["n"]
        conn.execute("UPDATE sessions SET sections = array_remove(sections, %s) "
                     "WHERE course_id = %s AND sections IS NOT NULL", (section, course_id))
    note = f"{short_section(section)}: {len(gone)} estudiantes quitados de la clase."
    if kept:
        note += f" {kept} se quedan porque ya tienen asistencia."
    return RedirectResponse(f"/course/{course_id}?msg={note}", status_code=303)


@app.post("/session/{session_id}/end")
def end_session(session_id: int, request: Request):
    with db() as conn:
        teacher = require_teacher(request, conn)
        own_session(conn, teacher, session_id)
        conn.execute("UPDATE sessions SET active = FALSE, ended_at = now() "
                     "WHERE id = %s AND active", (session_id,))
        # Nonces and tickets are worthless once the class is closed.
        conn.execute("DELETE FROM nonces WHERE session_id = %s", (session_id,))
        conn.execute("DELETE FROM scan_tickets WHERE session_id = %s", (session_id,))
    return RedirectResponse(f"/session/{session_id}/report", status_code=303)


def base_url(request: Request) -> str:
    proto = request.headers.get("x-forwarded-proto", request.url.scheme)
    host = request.headers.get("x-forwarded-host", request.headers.get("host", ""))
    return f"{proto}://{host}"


@app.get("/session/{session_id}/qr")
def attendance_qr(session_id: int, request: Request):
    """
    Mint a nonce and return the attendance QR. TEACHER ONLY: in version 1 this
    endpoint was public for a while, and a registered student could fetch a
    valid code from home. The QR must only exist where the professor projects it.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        if not session["active"]:
            return Response(status_code=410)
        nonce = secrets.token_hex(8)
        conn.execute("INSERT INTO nonces (nonce, session_id, expires_at) VALUES (%s, %s, %s)",
                     (nonce, session_id, utcnow() + timedelta(seconds=NONCE_VALID_SECONDS)))
    img = qrcode.make(f"{base_url(request)}/checkin?s={session_id}&n={nonce}")
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="image/png",
                             headers={"Cache-Control": "no-store"})


# ---------------------------------------------------------------
# CHECK-IN  (open to students: one QR, one tap)
# ---------------------------------------------------------------

def session_for_student(conn, session_id):
    return conn.execute("""
        SELECT s.*, c.name AS course_name, c.has_roster, t.name AS teacher_name,
               (c.has_roster AND NOT EXISTS (
                   SELECT 1 FROM enrollments en JOIN students st ON st.id = en.student_id
                   WHERE en.course_id = c.id AND st.username IS NOT NULL)) AS codes_only
        FROM sessions s JOIN courses c ON c.id = s.course_id
        JOIN teachers t ON t.id = c.teacher_id WHERE s.id = %s
    """, (session_id,)).fetchone()


def class_over_page():
    return student_page("<p class='bad-title'>Esta clase ya terminó</p>"
                        "<p>Si viniste a clase, avísale al profesor para que te marque a mano.</p>")


def rescan_page(reason="Ese QR ya venció."):
    return student_page(f"<p class='bad-title'>Vuelve a escanear</p><p>{e(reason)} "
                        "Apunta la cámara al QR que está ahora en la pantalla.</p>")


def accept_form(session, ticket, student=None, error="", ask_name=False, carne=""):
    msg = f"<div class='msg bad'>{e(error)}</div>" if error else ""
    head = (f"<div class='course'>{e(session['course_name'])}</div>"
            f"<p class='prof'>{e(session['teacher_name'])}</p>")
    if student:
        return student_page(f"""{head}{msg}
          <div class="name">{e(student['name'])}</div>
          <p class="carne">Carné {e(student['carne'])}</p>
          <form method="post" action="/checkin/accept">
            <input type="hidden" name="ticket" value="{e(ticket)}">
            <button class="btn big" type="submit">Aceptar</button>
          </form>""")
    name_field = "" if session["has_roster"] and not ask_name else """
            <label for="nombre">Nombre completo</label>
            <input id="nombre" name="nombre" required maxlength="80" autocomplete="name">"""
    # On October 26 the list came from Banner, which has student codes but no
    # usernames, and the form asked for "carné o usuario": many students typed
    # their username and waited as pending. When the list only has codes, the
    # form asks for the code.
    if session.get("codes_only"):
        what, label, field, hint = ("tu código de estudiante", "Código de estudiante", "numeric",
                                    "Es el número de tu carné de la universidad. No escribas tu usuario ni tu correo.")
    else:
        what, label, field, hint = ("tu carné o tu usuario", "Carné o usuario", "text",
                                    "Tu usuario es lo que va antes de @ en tu correo de la universidad.")
    return student_page(f"""{head}{msg}
      <p><strong>Primera vez en este celular.</strong> Escribe {what}
      una sola vez; las próximas clases solo tocarás Aceptar.</p>
      <form method="post" action="/checkin/accept">
        <input type="hidden" name="ticket" value="{e(ticket)}">
        <label for="carne">{label}</label>
        <input id="carne" name="carne" required inputmode="{field}" autocapitalize="none"
               autocomplete="off" value="{e(carne)}" autofocus>
        <p class="hint">{hint}</p>
        {name_field}
        <label style="display:flex;gap:10px;align-items:flex-start;margin-top:16px;font-weight:normal">
          <input type="checkbox" name="acepto" value="1" required style="width:auto;margin-top:4px">
          <span>Acepto la <a href="/privacidad" target="_blank">política de tratamiento de datos</a>:
          se guardan mi carné o usuario, mi nombre y la hora en que marco, solo para la asistencia.</span>
        </label>
        <p style="margin-top:16px"><button class="btn big" type="submit">Aceptar</button></p>
      </form>
      <p class="hint">Usa siempre este mismo navegador para escanear.</p>""")


def receipt_page(session, student, att, created):
    title = "Asistencia recibida" if created else "Ya tenías asistencia en esta clase"
    marked = local_time(att["marked_at"])
    return student_page(f"""
      <div class="okscreen"><div class="inner">
        <div class="check">✓</div>
        <p class="oktitle">{title}</p>
        <div class="name">{e(student['name'])}</div>
        <p class="okline">Carné {e(student['carne'])}</p>
        <p class="okline">{e(session['course_name'])} · {e(session['teacher_name'])}</p>
        <p class="okline">Marcada a las {marked}</p>
        <div class="clock" id="clock">--:--:--</div>
        <span class="live"><span class="dot"></span>hora actual, en vivo</span>
        <div class="receipt">Comprobante<br><b>{e(att['receipt'])}</b></div>
      </div></div>
      <script>
        // A live clock: a screenshot of this screen from another day, or
        // another minute, shows a time that does not match the room's.
        function tick() {{
          document.getElementById('clock').textContent =
            new Date().toLocaleTimeString('es-CO', {{ timeZone: 'America/Bogota', hour12: false }});
        }}
        tick(); setInterval(tick, 1000);
      </script>""")


def set_device_cookie(response: Response, device_id: str):
    response.set_cookie("device_id", device_id, max_age=31536000,
                        httponly=True, samesite="lax")


@app.get("/checkin", response_class=HTMLResponse)
def checkin(request: Request, s: int, n: str):
    device_id = request.cookies.get("device_id")
    new_cookie = not device_id
    if new_cookie:
        device_id = secrets.token_hex(16)
    ua = ua_of(request)

    with db() as conn:
        session = session_for_student(conn, s)
        if not session or not session["active"]:
            page = class_over_page()
        else:
            nonce = conn.execute("SELECT * FROM nonces WHERE nonce = %s AND session_id = %s",
                                 (n, s)).fetchone()
            if nonce is None or utcnow() > nonce["expires_at"]:
                # Counted, not flagged: how often students hit an expired code
                # is a usability number worth reporting. If the phone is known,
                # the report can say who it was and whether they marked later.
                known = None if new_cookie else conn.execute(
                    "SELECT id FROM students WHERE device_id = %s", (device_id,)).fetchone()
                log_event(conn, s, known["id"] if known else None, "qr_vencido",
                          ua_text(ua), False, ua)
                page = rescan_page()
            else:
                student = conn.execute("SELECT * FROM students WHERE device_id = %s",
                                       (device_id,)).fetchone()
                existing = None
                if student:
                    existing = conn.execute(
                        "SELECT * FROM attendance WHERE session_id = %s AND student_id = %s",
                        (s, student["id"])).fetchone()
                if existing:
                    page = receipt_page(session, student, existing, created=False)
                else:
                    ticket = secrets.token_urlsafe(16)
                    # Seconds since this QR was issued. A relayed QR (photo,
                    # video call) arrives older than one scanned in the room.
                    age = NONCE_VALID_SECONDS - (nonce["expires_at"] - utcnow()).total_seconds()
                    conn.execute("INSERT INTO scan_tickets (token, session_id, qr_age_s) "
                                 "VALUES (%s, %s, %s)", (ticket, s, round(age, 2)))
                    page = accept_form(session, ticket, student)

    response = HTMLResponse(page)
    if new_cookie:
        set_device_cookie(response, device_id)
    return response


@app.post("/checkin/accept", response_class=HTMLResponse)
def checkin_accept(request: Request, ticket: str = Form(...),
                   carne: str = Form(""), nombre: str = Form(""), acepto: str = Form("")):
    device_id = request.cookies.get("device_id")
    if not device_id:
        return student_page("<p class='bad-title'>Tu navegador no guardó la sesión</p>"
                            "<p>Revisa que no estés bloqueando cookies y vuelve a escanear.</p>")
    ua = ua_of(request)

    with db() as conn:
        t = conn.execute("SELECT * FROM scan_tickets WHERE token = %s FOR UPDATE",
                         (ticket,)).fetchone()
        if not t or t["used_at"]:
            return rescan_page("Ese escaneo ya se usó.")
        if utcnow() - t["scanned_at"] > timedelta(seconds=TICKET_SECONDS):
            return rescan_page("Pasó demasiado tiempo desde que escaneaste.")
        session = session_for_student(conn, t["session_id"])
        if not session or not session["active"]:
            return class_over_page()
        sid, course_id = session["id"], session["course_id"]

        student = conn.execute("SELECT * FROM students WHERE device_id = %s",
                               (device_id,)).fetchone()
        first_link = student is None

        if student is None:
            # First time in this browser: link it to the carne they type.
            carne = normalize_carne(carne)
            if not carne:
                return accept_form(session, ticket, error="Escribe tu carné.")
            if not acepto:
                # Nothing about this person is stored without their consent.
                return accept_form(session, ticket, carne=carne, ask_name=bool(nombre), error=(
                    "Para continuar, acepta la política de tratamiento de datos."))
            student = find_student(conn, carne)

            if session["has_roster"]:
                if student is None or not enrollment(conn, course_id, student["id"]):
                    # Version 1 locked 7 present students out. A student whose
                    # carne or username is not on the list is not refused: the
                    # first time they are asked to check it, the second time
                    # they are marked as pending and the professor decides.
                    nombre = re.sub(r"\s+", " ", nombre).strip()[:80]
                    if not nombre:
                        return accept_form(session, ticket, ask_name=True, carne=carne, error=(
                            "No encontramos ese carné o usuario en la lista de esta clase. "
                            "Revisa cómo lo escribiste. Si está bien, escribe también tu "
                            "nombre y quedas marcado para que el profesor lo revise."))
                    log_event(conn, sid, student["id"] if student else None,
                              "carne_no_en_lista", f"{carne} · {ua_text(ua)}", True, ua)
                    if student is None:
                        conn.execute("INSERT INTO students (carne, name) VALUES (%s, %s) "
                                     "ON CONFLICT (carne) DO NOTHING", (carne, nombre))
                        student = conn.execute("SELECT * FROM students WHERE carne = %s",
                                               (carne,)).fetchone()
                    conn.execute(
                        "INSERT INTO enrollments (course_id, student_id, status) "
                        "VALUES (%s, %s, 'pendiente') ON CONFLICT DO NOTHING",
                        (course_id, student["id"]))
            else:
                if student is None:
                    nombre = re.sub(r"\s+", " ", nombre).strip()[:80]
                    if not nombre:
                        return accept_form(session, ticket, error="Escribe tu nombre completo.")
                    conn.execute("INSERT INTO students (carne, name) VALUES (%s, %s) "
                                 "ON CONFLICT (carne) DO NOTHING", (carne, nombre))
                    student = conn.execute("SELECT * FROM students WHERE carne = %s",
                                           (carne,)).fetchone()
                conn.execute(
                    "INSERT INTO enrollments (course_id, student_id, status) "
                    "VALUES (%s, %s, 'pendiente') ON CONFLICT DO NOTHING",
                    (course_id, student["id"]))

            status = bind_device(conn, student, device_id, sid)
            if status == "taken":
                return student_page("<p class='bad-title'>Este navegador ya es de otro estudiante</p>"
                                    "<p>Un navegador solo puede marcar por una persona.</p>")
            student = conn.execute("SELECT * FROM students WHERE id = %s",
                                   (student["id"],)).fetchone()
            conn.execute("UPDATE students SET consent_at = COALESCE(consent_at, now()) "
                         "WHERE id = %s", (student["id"],))
            log_event(conn, sid, student["id"], "vinculacion", ua_text(ua), False, ua)
            if status == "moved":
                log_event(conn, sid, student["id"], "cambio_dispositivo",
                          f"ahora {ua_text(ua)}", True, ua)
            check_burst(conn, sid, ua)
        else:
            # A browser the system knows. Whatever carne might be in the form
            # is ignored: this browser marks for its owner and nobody else.
            if not enrollment(conn, course_id, student["id"]):
                # Same rule as for a new browser: a student who is not on the
                # list is not sent away. The check-in is recorded as pending
                # and flagged, and the professor decides.
                if session["has_roster"]:
                    log_event(conn, sid, student["id"], "carne_no_en_lista",
                              f"{student['carne']} · {ua_text(ua)}", True, ua)
                conn.execute(
                    "INSERT INTO enrollments (course_id, student_id, status) "
                    "VALUES (%s, %s, 'pendiente') ON CONFLICT DO NOTHING",
                    (course_id, student["id"]))

        if ua["in_app"]:
            seen = conn.execute(
                "SELECT 1 FROM events WHERE session_id = %s AND student_id = %s "
                "AND kind = 'desde_app'", (sid, student["id"])).fetchone()
            if not seen:
                log_event(conn, sid, student["id"], "desde_app", ua["in_app"], True, ua)

        att, created = record_attendance(conn, session, student, "qr", ua)
        conn.execute("UPDATE scan_tickets SET used_at = now() WHERE token = %s", (ticket,))
        if created:
            conn.execute(
                "UPDATE attendance SET qr_age_s = %s, first_link = %s, "
                "accept_s = EXTRACT(EPOCH FROM now() - %s::timestamptz) WHERE id = %s",
                (t["qr_age_s"], first_link, t["scanned_at"], att["id"]))
        return receipt_page(session, student, att, created)


# ---------------------------------------------------------------
# REPORTS
# ---------------------------------------------------------------

def build_report(conn, session, late_minutes: int):
    """
    One row per student who belongs to the class or who was marked in it.
    The same rows feed the screen, the PDF, Excel and CSV, so the four can
    never disagree.
    """
    groups = session["sections"]
    students = conn.execute("""
        SELECT st.id, st.carne, st.name, st.last_name, en.status, en.section FROM students st
        LEFT JOIN enrollments en ON en.student_id = st.id AND en.course_id = %s
        WHERE (en.course_id IS NOT NULL
               AND (%s::text[] IS NULL OR en.section = ANY(%s::text[])))
           OR st.id IN (SELECT student_id FROM attendance WHERE session_id = %s)
        ORDER BY st.name
    """, (session["course_id"], groups, groups, session["id"])).fetchall()
    attendance = {a["student_id"]: a for a in conn.execute(
        "SELECT * FROM attendance WHERE session_id = %s", (session["id"],)).fetchall()}
    moved = {r["student_id"] for r in conn.execute(
        "SELECT DISTINCT student_id FROM device_changes WHERE session_id = %s",
        (session["id"],)).fetchall()}
    rejected = {r["student_id"] for r in conn.execute(
        "SELECT DISTINCT student_id FROM events WHERE session_id = %s "
        "AND decision = 'rechazado'", (session["id"],)).fetchall()}
    notes = {}
    for ev in conn.execute(
            "SELECT * FROM events WHERE session_id = %s AND review AND student_id IS NOT NULL "
            "ORDER BY created_at", (session["id"],)).fetchall():
        label = EVENT_LABELS.get(ev["kind"], ev["kind"])
        if ev["decision"]:
            label += f" ({ev['decision']} por {ev['decided_by']})"
        notes.setdefault(ev["student_id"], []).append(label)

    rows = []
    for st in students:
        a = attendance.get(st["id"])
        minutes = None
        if a:
            minutes = int((a["marked_at"] - session["started_at"]).total_seconds() // 60)
            if st["id"] in rejected:
                state = "Rechazado"
            elif a["method"] == "manual":
                state = "Manual"
            elif minutes > late_minutes:
                state = "Tarde"
            else:
                state = "Presente"
        else:
            state = "Ausente"
        observation = list(dict.fromkeys(notes.get(st["id"], [])))
        if st["status"] == "pendiente":
            observation.insert(0, "Se unió marcando; falta aprobarlo")
        if groups and st["section"] and st["section"] not in groups:
            observation.insert(0, f"Es de otro grupo ({short_section(st['section'])})")
        rows.append({
            "carne": st["carne"], "nombre": st["name"],
            "grupo": short_section(st["section"]) if st["section"] else "",
            "estado": state,
            "hora": local_time(a["marked_at"]) if a else "",
            "retraso_min": minutes if a else "",
            "metodo": ("QR" if a["method"] == "qr" else "Manual") if a else "",
            "navegador": (a["browser"] or "") if a else "",
            "sistema": (a["os"] or "") if a else "",
            "app": (a["in_app"] or "") if a else "",
            "celular": f"Celular #{a['device_no']}" if a and a["device_no"] else "",
            "cambio_dispositivo": "Sí" if st["id"] in moved else "No",
            "comprobante": a["receipt"] if a else "",
            "observacion": "; ".join(observation),
            "_apellido": _plain(f"{st['last_name']} {st['name']}" if st["last_name"] else st["name"]),
            "_llegada": a["marked_at"].timestamp() if a else None,
        })
    return rows


REPORT_COLUMNS = [
    ("carne", "Carné"), ("nombre", "Nombre"), ("grupo", "Grupo"), ("estado", "Estado"),
    ("hora", "Hora"),
    ("retraso_min", "Minutos desde el inicio"), ("metodo", "Método"),
    ("navegador", "Navegador"), ("sistema", "Sistema"), ("app", "App de origen"),
    ("celular", "Celular"), ("cambio_dispositivo", "Cambio de dispositivo"),
    ("comprobante", "Comprobante"), ("observacion", "Observación"),
]


def report_context(request: Request, session_id: int, late: int):
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        rows = build_report(conn, session, late)
        alerts = conn.execute("""
            SELECT ev.*, st.name FROM events ev LEFT JOIN students st ON st.id = ev.student_id
            WHERE ev.session_id = %s AND ev.review
            ORDER BY (ev.kind = 'rafaga'), ev.created_at
        """, (session_id,)).fetchall()
        # Who scanned an expired QR, when the phone was known, and whether they
        # managed to mark afterwards. New phones cannot be named.
        expired = conn.execute("""
            SELECT ev.created_at, ev.student_id, st.name,
                   EXISTS (SELECT 1 FROM attendance a WHERE a.session_id = ev.session_id
                           AND a.student_id = ev.student_id) AS marked
            FROM events ev LEFT JOIN students st ON st.id = ev.student_id
            WHERE ev.session_id = %s AND ev.kind = 'qr_vencido' ORDER BY ev.created_at
        """, (session_id,)).fetchall()
    return teacher, session, rows, alerts, expired


def file_name(session, ext):
    stem = re.sub(r"[^A-Za-z0-9]+", "_", _plain(session["course_name"])).strip("_") or "clase"
    day = session["started_at"].astimezone(LOCAL_TZ).strftime("%Y-%m-%d")
    return f"asistencia_{stem}_{day}.{ext}"


@app.get("/session/{session_id}/report", response_class=HTMLResponse)
def report_page(session_id: int, request: Request, tarde: int = DEFAULT_LATE_MINUTES):
    teacher, session, rows, alerts, expired = report_context(request, session_id, tarde)
    count = lambda state: sum(1 for r in rows if r["estado"] == state)
    present = count("Presente") + count("Tarde") + count("Manual")
    head = "".join(f"<th>{label}</th>" for _, label in REPORT_COLUMNS)
    def sort_attrs(r):
        arrival = "" if r["_llegada"] is None else f"{r['_llegada']:.3f}"
        return (f" data-apellido='{e(r['_apellido'])}' data-nombre='{e(_plain(r['nombre']))}' "
                f"data-carne='{e(r['carne'])}' data-llegada='{arrival}' data-estado='{e(r['estado'])}'")
    body = "".join(
        f"<tr{sort_attrs(r)}>" + "".join(
            f"<td>{e(r[key])}</td>" if key != "estado" else
            f"<td><span class='pill {'bad' if r[key] in ('Ausente', 'Rechazado') else 'warn' if r[key] in ('Tarde', 'Manual') else 'ok'}'>{e(r[key])}</span></td>"
            for key, _ in REPORT_COLUMNS) + "</tr>"
        for r in rows)
    def alert_li(a):
        return (f"<li>{local_time(a['created_at'])} · {e(a['name'] or 'Varios estudiantes')} · "
                f"{e(EVENT_LABELS.get(a['kind'], a['kind']))}"
                f"{' (' + e(a['detail']) + ')' if a['detail'] else ''}</li>")
    # Alerts about one student first; the general notice about many new phones
    # in a row goes after them, so it does not bury the ones that need a decision.
    alert_items = "".join(alert_li(a) for a in alerts if a["kind"] != "rafaga")
    bursts = "".join(alert_li(a) for a in alerts if a["kind"] == "rafaga")
    if bursts:
        alert_items += ("</ul><p class='hint' style='margin:10px 0 4px'>Avisos generales "
                        "(normales la primera vez que un grupo usa la app):</p><ul>" + bursts)
    def expired_li(x):
        if not x["student_id"]:
            who = "Celular nuevo (todavía sin nombre)"
            status = ""
        elif x["marked"]:
            who, status = e(x["name"]), " · <span class='pill ok'>después marcó</span>"
        else:
            who = e(x["name"])
            status = (f" · <span class='pill bad'>no marcó</span> "
                      f"<button class='btn small noprint' type='button' "
                      f"onclick=\"markLate({x['student_id']}, this)\">Marcar presente</button>")
        return f"<li>{local_time(x['created_at'])} · {who}{status}</li>"
    expired_html = ("<ul>" + "".join(expired_li(x) for x in expired) + "</ul>" if expired
                    else "<p class='muted'>Nadie escaneó un QR vencido.</p>")
    groups_line = (f"<p>Grupos de esta sesión: {e(', '.join(short_section(g) for g in session['sections']))}</p>"
                   if session["sections"] else "")
    with db() as conn:
        course_groups = course_sections(conn, session["course_id"])
    if len(course_groups) > 1:
        # A session taken before the list existed has no groups, so the report
        # would count every group of the list. The professor sets them here.
        mine = set(session["sections"] or [])
        boxes = "".join(
            f"<label style='font-weight:normal;display:inline-flex;gap:6px;align-items:center;margin-right:14px'>"
            f"<input type='checkbox' name='sections' value='{e(g)}' {'checked' if g in mine else ''} style='width:auto'>"
            f"{e(short_section(g))} <span class='muted'>({n})</span></label>"
            for g, n in course_groups.items())
        groups_line += f"""
      <details class="noprint" {'open' if not mine else ''} style="margin-bottom:12px">
        <summary>¿Qué grupos estaban en esta sesión?</summary>
        <form method="post" action="/session/{session_id}/groups" class="card" style="margin-top:8px">
          <p class="hint">Si tomaste asistencia antes de subir la lista, marca aquí los grupos que
          estaban en el salón. El reporte cuenta solo esos grupos; nadie que marcó asistencia se pierde.</p>
          <div>{boxes}</div>
          <p style="margin-top:10px"><button class="btn small" type="submit">Guardar grupos</button></p>
        </form>
      </details>"""
    alerts_html = (f"<ul>{alert_items}</ul>" if alert_items
                   else "<p class='muted'>Nada que revisar en esta clase.</p>")
    ended = local(session["ended_at"]) if session["ended_at"] else "en curso"
    live_link = (f"<a class='btn' href='/session/{session_id}/live'>Volver a la pantalla en vivo</a>"
                 if session["active"] else "")
    institution = f"{e(teacher['institution'])} · " if teacher["institution"] else ""

    return teacher_page(f"Reporte {session['course_name']}", f"""
      <style>
        .stats {{ display:flex; gap:10px; flex-wrap:wrap; margin:10px 0 18px; }}
        .stat {{ background:var(--card); border:1px solid var(--line); border-radius:10px; padding:10px 14px;
                 font:inherit; color:inherit; text-align:left; text-decoration:none; }}
        button.stat, a.stat {{ cursor:pointer; }}
        .stat.on {{ outline:2px solid var(--ok); }}
        .stat b {{ display:block; font-size:26px; }}
        .sign {{ margin-top:48px; display:flex; gap:40px; flex-wrap:wrap; }}
        .sign div {{ border-top:1px solid var(--ink); padding-top:6px; min-width:240px; }}
        @media print {{
          header.bar, .noprint {{ display:none !important; }}
          body {{ background:#fff; }} .wrap {{ max-width:none; padding:0; }}
          table {{ font-size:10px; }} th, td {{ padding:4px 5px; }}
          .table-wrap {{ overflow:visible; }}
        }}
      </style>
      <div class="noprint" style="display:flex;gap:8px;flex-wrap:wrap;margin-bottom:12px">
        <button class="btn" type="button" onclick="window.print()">Imprimir o guardar PDF</button>
        <a class="btn light" href="/session/{session_id}/export.xlsx?tarde={tarde}">Descargar Excel</a>
        <a class="btn light" href="/session/{session_id}/export.csv?tarde={tarde}">Descargar CSV</a>
        <a class="btn light" href="/session/{session_id}/timing.csv">Tiempos (anónimo)</a>
        {live_link}
        <a class="btn grey" href="/course/{session['course_id']}">Volver a la clase</a>
      </div>
      <p class="muted" style="margin:0">{institution}Reporte de asistencia</p>
      <h1 style="margin-top:4px">{e(session['course_name'])}</h1>
      <p>Profesor: {e(teacher['name'])} · Inicio: {local(session['started_at'])} · Fin: {ended} · Sesión #{session_id}</p>
      {groups_line}
      <div class="stats">
        <button type="button" class="stat" onclick="filterReport('', this)"><b>{present} de {len(rows)}</b>asistieron · ver todos</button>
        <button type="button" class="stat" onclick="filterReport('Ausente', this)"><b>{count('Ausente')}</b>ausentes</button>
        <button type="button" class="stat" onclick="filterReport('Tarde', this)"><b>{count('Tarde')}</b>llegaron después de {tarde} min</button>
        <button type="button" class="stat" onclick="filterReport('Manual', this)"><b>{count('Manual')}</b>marcados a mano</button>
        <a class="stat" href="#por-revisar"><b>{len(alerts)}</b>por revisar</a>
        <a class="stat" href="#qr-vencido" title="Alguien escaneó un QR viejo y la app le pidió escanear otra vez."><b>{len(expired)}</b>escaneos de QR vencido</a>
      </div>
      <h2 id="por-revisar">Por revisar</h2>
      {alerts_html}
      <p class="hint">Estas señales son pistas, no pruebas. Si algo no cuadra, pregúntale al estudiante.</p>
      <h2 id="qr-vencido">Escanearon un QR vencido</h2>
      {expired_html}
      <p class="hint">La app les pidió escanear otra vez. Si alguien no alcanzó a marcar y sí estaba, márcalo aquí.</p>
      <p id="lateMsg" class="hint"></p>
      <script>
      async function markLate(id, btn) {{
        const f = new FormData(); f.append('student_id', id);
        const r = await fetch('/session/{session_id}/manual', {{ method: 'POST', body: f }});
        const d = await r.json().catch(() => ({{ message: 'No se pudo marcar.' }}));
        document.getElementById('lateMsg').textContent = d.message;
        if (r.ok) btn.replaceWith(document.createTextNode('marcado a mano'));
      }}
      </script>
      <form class="noprint" method="get" style="display:flex;gap:8px;align-items:center;max-width:420px;margin-bottom:12px">
        <label for="tarde" style="margin:0;white-space:nowrap">Tarde después de</label>
        <input id="tarde" name="tarde" type="number" min="0" max="240" value="{tarde}" style="width:90px">
        <span>min</span><button class="btn small" type="submit">Aplicar</button>
      </form>
      <form class="noprint" style="display:flex;gap:8px;align-items:center;max-width:460px;margin-bottom:10px"
            onsubmit="return false">
        <label for="orden" style="margin:0;white-space:nowrap">Ordenar por</label>
        <select id="orden" onchange="sortReport(this.value)">
          <option value="apellido">Apellido, A a Z</option>
          <option value="apellido-desc">Apellido, Z a A</option>
          <option value="carne">Código</option>
          <option value="llegada">Llegada: primero en marcar</option>
          <option value="llegada-desc">Llegada: último en marcar</option>
          <option value="estado">Estado</option>
        </select>
      </form>
      <div class="table-wrap"><table id="report"><tr>{head}</tr>{body}</table></div>
      <script>
      function sortReport(by) {{
        const desc = by.endsWith('-desc'); const key = by.replace('-desc', '');
        const table = document.getElementById('report');
        const rows = Array.from(table.querySelectorAll('tr')).slice(1);
        rows.sort((a, b) => {{
          let x = a.dataset[key] || '', y = b.dataset[key] || '';
          if (key === 'llegada') {{
            if (!x) return 1; if (!y) return -1;     // absent students always last
            return desc ? y - x : x - y;
          }}
          return desc ? y.localeCompare(x, 'es', {{numeric: true}}) : x.localeCompare(y, 'es', {{numeric: true}});
        }});
        rows.forEach(r => table.appendChild(r));
      }}
      function filterReport(state, tile) {{
        document.querySelectorAll('.stat').forEach(t => t.classList.remove('on'));
        if (state) tile.classList.add('on');
        document.querySelectorAll('#report tr').forEach((r, i) => {{
          if (i) r.style.display = (!state || r.dataset.estado === state) ? '' : 'none';
        }});
      }}
      sortReport('apellido');
      </script>
      <div class="sign"><div>Firma del profesor</div><div>Fecha</div></div>
    """, teacher)


def sheet_safe(value):
    """
    A name typed by a student, like =HYPERLINK("http://..."), would run as a
    formula when the professor opens the report in Excel (CSV injection).
    Text that starts with = + - or @ gets an apostrophe in front, so Excel
    shows it as text. Numbers are left alone.
    """
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


@app.get("/session/{session_id}/export.csv")
def export_csv(session_id: int, request: Request, tarde: int = DEFAULT_LATE_MINUTES):
    _, session, rows, _, _ = report_context(request, session_id, tarde)
    buffer = io.StringIO()
    # The csv module quotes any value that contains a comma; version 1 did not,
    # and a name like "Pérez, Juan" shifted every column after it.
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in REPORT_COLUMNS])
    for r in rows:
        writer.writerow([sheet_safe(r[key]) for key, _ in REPORT_COLUMNS])
    data = ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM: Excel reads the accents
    return Response(data, media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={file_name(session, 'csv')}"})


TIMING_COLUMNS = ["n", "qr_age_s", "accept_s", "first_link", "browser", "os", "in_app"]


@app.get("/session/{session_id}/timing.csv")
def export_timing(session_id: int, request: Request):
    """
    One row per QR check-in, in the order they happened, with no name, carne
    or receipt: only how old the QR was when scanned and how long the student
    took to tap Aceptar. Anonymous by construction, so it can go straight into
    the repository for the relay analysis.
    """
    with db() as conn:
        teacher = require_teacher(request, conn)
        session = own_session(conn, teacher, session_id)
        rows = conn.execute(
            "SELECT qr_age_s, accept_s, first_link, browser, os, in_app FROM attendance "
            "WHERE session_id = %s AND method = 'qr' ORDER BY marked_at", (session_id,)).fetchall()
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TIMING_COLUMNS)
    for i, r in enumerate(rows, start=1):
        writer.writerow([i, "" if r["qr_age_s"] is None else round(r["qr_age_s"], 2),
                         "" if r["accept_s"] is None else round(float(r["accept_s"]), 2),
                         "" if r["first_link"] is None else int(r["first_link"]),
                         r["browser"] or "", r["os"] or "", r["in_app"] or ""])
    return Response(buffer.getvalue().encode("utf-8"), media_type="text/csv",
                    headers={"Content-Disposition":
                             f"attachment; filename=tiempos_sesion_{session['id']}.csv"})


@app.get("/session/{session_id}/export.xlsx")
def export_xlsx(session_id: int, request: Request, tarde: int = DEFAULT_LATE_MINUTES):
    import openpyxl
    from openpyxl.styles import Font
    _, session, rows, _, _ = report_context(request, session_id, tarde)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asistencia"
    ws.append([label for _, label in REPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([sheet_safe(r[key]) for key, _ in REPORT_COLUMNS])
    widths = [14, 30, 12, 11, 10, 12, 9, 16, 10, 14, 12, 12, 16, 50]
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(buffer.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={file_name(session, 'xlsx')}"})
